import os
import json
import logging
import re
import unicodedata
from datetime import datetime, timedelta, timezone
from typing import Optional, List
from urllib.parse import quote
from fastapi import FastAPI, Header, HTTPException, Request, BackgroundTasks, Form, UploadFile, File
from fastapi.encoders import jsonable_encoder
from fastapi.responses import HTMLResponse, RedirectResponse, FileResponse, Response
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from fastapi.websockets import WebSocket, WebSocketDisconnect
from sqlalchemy import text
from sqlalchemy.exc import ProgrammingError
from sqlmodel import Session, create_engine, select, func, SQLModel, desc
from models import Usuario, Impressora, Impressao, Licenca, Chamado, ChamadoInteracao, TipoProblema, EquipamentoTI, MarmitaCardapio, MarmitaPedido, EstudoSessao, EstudoPlano, EstudoCronometro, EstudoRaioXBanca
import license_service
import api_printers
import api_reports
import api_users_monitor
import api_dashboard
import api_users_mgmt
import api_settings
import api_auth
import api_tickets
import api_finance
import api_avisos
import api_status
import api_equipamentos_ti
import api_cofre
import api_whatsapp
import api_tenants
from database import engine
from dotenv import load_dotenv
from websocket_manager import manager
from auth_utils import get_signed_cookie, normalize_permissions, verify_study_access_token
from atendimento_tenant_context import resolve_context

from slowapi import Limiter, _rate_limit_exceeded_handler
from slowapi.util import get_remote_address
from slowapi.errors import RateLimitExceeded

load_dotenv()

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
logger = logging.getLogger(__name__)

# FunÃ§Ã£o para carregar configuraÃ§Ãµes originais do PHP
def load_app_config():
    config_path = os.path.join(BASE_DIR, "public_html", "includes", "config.json")
    default_config = {
        "APP_NAME": "PrintDash",
        "APP_COLOR": "#6366f1",
        "APP_LOGO_URL": ""
    }
    if os.path.exists(config_path):
        try:
            with open(config_path, "r", encoding="utf-8-sig") as f:
                return json.load(f)
        except (OSError, json.JSONDecodeError) as exc:
            logger.error("Não foi possível carregar %s: %s", config_path, exc)
    return default_config

APP_CONFIG = load_app_config()

# engine agora vem do database.py
app = FastAPI(title=APP_CONFIG.get("APP_NAME", "PrintDash"))

@app.middleware("http")
async def restrict_study_accounts(request: Request, call_next):
    """Keep study-only accounts inside the independent study area."""
    path = request.url.path
    if path.startswith("/estudos") and path != "/estudos/login":
        study_user_id = verify_study_access_token(request.cookies.get("study_access"))
        if not study_user_id or request.cookies.get("user_id") != study_user_id:
            return RedirectResponse(url="/estudos/login", status_code=303)

    if request.cookies.get("user_role") == "estudante":
        allowed = (
            path.startswith("/estudos")
            or path.startswith("/api/auth")
            or path.startswith("/static")
            or path.startswith("/uploads")
        )
        if not allowed:
            return RedirectResponse(url="/estudos", status_code=303)
    return await call_next(request)

from fastapi.middleware.cors import CORSMiddleware
app.add_middleware(
    CORSMiddleware,
    allow_origin_regex=".*",
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

limiter = Limiter(key_func=get_remote_address)
app.state.limiter = limiter
app.add_exception_handler(RateLimitExceeded, _rate_limit_exceeded_handler)

# Templates
templates = Jinja2Templates(directory=os.path.join(BASE_DIR, "templates"))
def from_json_filter(value):
    try:
        if not value: return []
        return json.loads(value)
    except:
        return []
templates.env.filters["from_json"] = from_json_filter

@app.websocket("/ws")
async def websocket_endpoint(websocket: WebSocket):
    await manager.connect(websocket)
    try:
        while True:
            data = await websocket.receive_text()
            # Respond to ping to keep connection alive
            if data == "ping":
                await websocket.send_text("pong")
    except WebSocketDisconnect:
        manager.disconnect(websocket)


def include_router_if_available(module, module_name: str):
    """Inclui um router apenas quando o mÃ³dulo realmente expÃµe `router`."""
    router = getattr(module, "router", None)
    if router is None:
        print(f"[STARTUP] Ignorando {module_name}: nenhum router exportado.")
        return

    app.include_router(router)
    print(f"[STARTUP] Router carregado: {module_name}")


def ensure_schema_compatibility():
    """Aplica ajustes pequenos no schema sem apagar dados existentes."""
    from sqlalchemy import text

    statements = [
        "ALTER TABLE impressoras ADD COLUMN IF NOT EXISTS modelo VARCHAR(120)",
        "ALTER TABLE usuarios ADD COLUMN IF NOT EXISTS empresa_nome VARCHAR(255)",
        "ALTER TABLE usuarios ADD COLUMN IF NOT EXISTS empresa_logo_url VARCHAR(500)",
        "ALTER TABLE usuarios ADD COLUMN IF NOT EXISTS foto_perfil_url VARCHAR(500)",
        "ALTER TABLE usuarios ADD COLUMN IF NOT EXISTS last_seen TIMESTAMP NULL",
        "ALTER TABLE usuarios DROP CONSTRAINT IF EXISTS usuarios_role_check",
        "ALTER TABLE usuarios ADD CONSTRAINT usuarios_role_check CHECK (role IN ('admin', 'operator', 'estudante', 'gestor', 'gerente', 'supervisor', 'recepcao', 'recepcionista', 'cliente_atendimento'))",
        "ALTER TABLE impressoras ADD COLUMN IF NOT EXISTS localizacao VARCHAR(120)",
        "ALTER TABLE tipos_problemas ADD COLUMN IF NOT EXISTS sla_horas INTEGER DEFAULT 24",
        "ALTER TABLE equipamentos_ti ADD COLUMN IF NOT EXISTS modelo_monitor VARCHAR(120)",
        "ALTER TABLE equipamentos_ti ADD COLUMN IF NOT EXISTS quantidade_monitores INTEGER",
        "ALTER TABLE equipamentos_ti ADD COLUMN IF NOT EXISTS nome_computador VARCHAR(120)",
        "ALTER TABLE equipamentos_ti ADD COLUMN IF NOT EXISTS data_registro TIMESTAMP NULL",
        "ALTER TABLE chamados ADD COLUMN IF NOT EXISTS ip_address TEXT",
        "ALTER TABLE chamados ADD COLUMN IF NOT EXISTS user_agent TEXT",
        "ALTER TABLE chamados ADD COLUMN IF NOT EXISTS unread_admin INTEGER DEFAULT 0",
        "ALTER TABLE chamados ADD COLUMN IF NOT EXISTS unread_user INTEGER DEFAULT 0",
        "ALTER TABLE chamados ADD COLUMN IF NOT EXISTS vencimento_sla TIMESTAMP NULL",
        "ALTER TABLE chamados ADD COLUMN IF NOT EXISTS avaliacao_estrelas INTEGER",
        "ALTER TABLE chamados ADD COLUMN IF NOT EXISTS avaliacao_comentario TEXT",
        "ALTER TABLE cofre_senhas ADD COLUMN IF NOT EXISTS setor VARCHAR(80) DEFAULT 'TI'",
        "ALTER TABLE chamados ADD COLUMN IF NOT EXISTS whatsapp_cliente VARCHAR(80)",
        "ALTER TABLE chamados ADD COLUMN IF NOT EXISTS whatsapp_instance VARCHAR(120)",
        "ALTER TABLE chamados ADD COLUMN IF NOT EXISTS origem VARCHAR(50) DEFAULT 'Web'",
        "ALTER TABLE chamados ADD COLUMN IF NOT EXISTS assigned_user VARCHAR(120)",
        "ALTER TABLE estudos_planos ADD COLUMN IF NOT EXISTS dia_semana VARCHAR(20)",
        "ALTER TABLE estudos_planos ADD COLUMN IF NOT EXISTS slot INTEGER",
        "ALTER TABLE estudos_planos ADD COLUMN IF NOT EXISTS horario_inicio VARCHAR(10)",
        "ALTER TABLE estudos_planos ADD COLUMN IF NOT EXISTS horario_fim VARCHAR(10)",
        "ALTER TABLE estudos_sessoes ADD COLUMN IF NOT EXISTS tipo_registro VARCHAR(20) DEFAULT 'estudo'",
        "ALTER TABLE estudos_sessoes ADD COLUMN IF NOT EXISTS nome_simulado VARCHAR(255)",
        "ALTER TABLE chamados ADD COLUMN IF NOT EXISTS visivel_suporte BOOLEAN DEFAULT TRUE",
        "ALTER TABLE chamados_interacoes ADD COLUMN IF NOT EXISTS whatsapp_status VARCHAR(20) DEFAULT 'sent'",
        "ALTER TABLE empresas_atendimento ADD COLUMN IF NOT EXISTS limite_usuarios_custom INTEGER",
        "ALTER TABLE empresas_atendimento ADD COLUMN IF NOT EXISTS limite_instancias_whatsapp_custom INTEGER",
        "ALTER TABLE empresas_atendimento ADD COLUMN IF NOT EXISTS logo_url VARCHAR(500)",
        "ALTER TABLE empresas_atendimento ADD COLUMN IF NOT EXISTS cor VARCHAR(20)",
        "ALTER TABLE status_empresa ADD COLUMN IF NOT EXISTS empresa_atendimento_id INTEGER",
        "ALTER TABLE status_empresa ADD COLUMN IF NOT EXISTS whatsapp_instance VARCHAR(120)",
        "ALTER TABLE impressoes DROP CONSTRAINT IF EXISTS impressoes_tipo_impressao_check",
        "ALTER TABLE impressoes ADD CONSTRAINT impressoes_tipo_impressao_check CHECK (tipo_impressao IN ('Preto e Branco', 'Colorida', 'Digital', 'P&B', 'Colorido'))",
        """CREATE TABLE IF NOT EXISTS setores_atendimento (
            id SERIAL PRIMARY KEY,
            empresa_atendimento_id INTEGER NOT NULL REFERENCES empresas_atendimento(id) ON DELETE CASCADE,
            nome VARCHAR(120) NOT NULL,
            ativo BOOLEAN NOT NULL DEFAULT TRUE,
            UNIQUE (empresa_atendimento_id, nome)
        )""",
        """CREATE TABLE IF NOT EXISTS usuarios_setores_atendimento (
            usuario_id INTEGER NOT NULL REFERENCES usuarios(id) ON DELETE CASCADE,
            setor_id INTEGER NOT NULL REFERENCES setores_atendimento(id) ON DELETE CASCADE,
            papel VARCHAR(30) NOT NULL DEFAULT 'operador',
            PRIMARY KEY (usuario_id, setor_id)
        )""",
        """CREATE TABLE IF NOT EXISTS instancias_setores_atendimento (
            instancia_id INTEGER NOT NULL REFERENCES instancias_whatsapp_tenant(id) ON DELETE CASCADE,
            setor_id INTEGER NOT NULL REFERENCES setores_atendimento(id) ON DELETE CASCADE,
            PRIMARY KEY (instancia_id, setor_id)
        )""",
    ]

    # Cada ajuste usa sua própria transação. Assim, uma alteração opcional
    # que falhe não desfaz as colunas essenciais de chamados/WhatsApp.
    for statement in statements:
        try:
            with engine.begin() as conn:
                conn.execute(text(statement))
        except Exception as exc:
            print(f"[SCHEMA] Ajuste ignorado: {statement[:90]}... ({exc})")


def ensure_whatsapp_columns():
    statements = [
        "ALTER TABLE chamados ADD COLUMN IF NOT EXISTS whatsapp_cliente VARCHAR(80)",
        "ALTER TABLE chamados ADD COLUMN IF NOT EXISTS whatsapp_instance VARCHAR(120)"
    ]
    with engine.begin() as conn:
        for statement in statements:
            conn.execute(text(statement))


@app.get("/chamados", response_class=HTMLResponse)
async def view_chamados(request: Request, success: str = None, error: str = None):
    print("[FORCE_CHAMADOS_ROUTE] Entering function")
    try:
        user_id = request.cookies.get("user_id")
        if not user_id: return RedirectResponse(url="/login")
        
        user_role = request.cookies.get("user_role")
        user_name = request.cookies.get("user_name")
        
        with Session(engine) as session:
            impressoras = session.exec(select(Impressora)).all()
            atalhos = session.exec(select(TipoProblema)).all()
            tenant_context = resolve_context(request)
            tenant_conditions = []
            if tenant_context.enabled and tenant_context.tenant_id is not None:
                tenant_conditions.append(Chamado.empresa_atendimento_id == tenant_context.tenant_id)
            
            hoje = datetime.now().date()
            amanha = hoje + timedelta(days=1)
            quatro_horas_atras = datetime.now(timezone.utc) - timedelta(hours=4)

            def como_utc(valor):
                if valor is None:
                    return None
                if valor.tzinfo is None:
                    # Registros antigos foram gravados sem fuso; no servidor eles representam UTC.
                    return valor.replace(tzinfo=timezone.utc)
                return valor.astimezone(timezone.utc)

            try:
                if user_role == "admin":
                    chamados = session.exec(select(Chamado).where(Chamado.visivel_suporte == True, *tenant_conditions).order_by(desc(Chamado.data_abertura))).all()
                else:
                    chamados = session.exec(select(Chamado).where(Chamado.usuario == user_name, Chamado.visivel_suporte == True, *tenant_conditions).order_by(desc(Chamado.data_abertura))).all()
            except ProgrammingError as pe:
                if "whatsapp_instance" in str(pe) or "whatsapp_cliente" in str(pe):
                    print(f"[STARTUP] Detectado erro de schema WhatsApp em chamados: {pe}. Aplicando correção de schema.")
                    # A consulta que falhou deixa a transação SQLAlchemy abortada.
                    # Sem rollback, a nova consulta falha mesmo após criar as colunas.
                    session.rollback()
                    ensure_whatsapp_columns()
                    if user_role == "admin":
                        chamados = session.exec(select(Chamado).where(Chamado.visivel_suporte == True, *tenant_conditions).order_by(desc(Chamado.data_abertura))).all()
                    else:
                        chamados = session.exec(select(Chamado).where(Chamado.usuario == user_name, Chamado.visivel_suporte == True, *tenant_conditions).order_by(desc(Chamado.data_abertura))).all()
                else:
                    raise

            total_tickets = len(chamados)
            abertos = sum(1 for c in chamados if c.status == "Aberto")
            atendimento = sum(1 for c in chamados if c.status == "Em Atendimento")
            resolvidos = sum(1 for c in chamados if c.status == "Resolvido")
            resolvidos_hoje = sum(1 for c in chamados if c.status == "Resolvido" and c.data_resolucao and hoje <= c.data_resolucao.date() < amanha)
            urgentes = sum(
                1 for c in chamados
                if c.status == "Aberto"
                and (data_abertura := como_utc(getattr(c, "data_abertura", None))) is not None
                and data_abertura < quatro_horas_atras
            )

            saude_val = round((resolvidos / total_tickets * 100)) if total_tickets > 0 else 100

            msg = ""
            t_msg = ""
            if success == "1":
                msg = "Protocolo registrado com sucesso!"
                t_msg = "success"
            elif success == "resolved":
                msg = "Chamado encerrado!"
                t_msg = "success"
            elif error == "duplicado":
                msg = "SolicitaÃ§Ã£o jÃ¡ recebida."
                t_msg = "warning"

            render_context = {
                "request": request,
                "config": APP_CONFIG,
                "is_admin": user_role == "admin",
                "impressoras": [i.dict() for i in impressoras],
                "atalhos": [a.dict() for a in atalhos],
                "message": msg,
                "tipo_msg": t_msg,
                "estatisticas_tickets": {
                    "abertos": int(abertos),
                    "atendimento": int(atendimento),
                    "resolvidos": int(resolvidos),
                    "resolvidos_hoje": int(resolvidos_hoje),
                    "urgentes": int(urgentes),
                    "total": int(total_tickets),
                    "saude": int(saude_val)
                },
                "chamados": [c.dict() for c in chamados]
            }
            return templates.TemplateResponse(request, "chamados.html", render_context)
    except Exception as e:
        import traceback
        err = traceback.format_exc()
        print(f"CRITICAL ERROR IN VIEW_CHAMADOS:\n{err}")
        return HTMLResponse(content=f"<h1>DEBUG ERROR</h1><pre>{err}</pre>", status_code=500)

include_router_if_available(api_printers, "api_printers")
include_router_if_available(api_reports, "api_reports")
include_router_if_available(api_dashboard, "api_dashboard")
include_router_if_available(api_tickets, "api_tickets")
include_router_if_available(api_users_monitor, "api_users_monitor")
include_router_if_available(api_users_mgmt, "api_users_mgmt")
include_router_if_available(api_settings, "api_settings")
include_router_if_available(api_auth, "api_auth")
include_router_if_available(api_finance, "api_finance")
include_router_if_available(api_avisos, "api_avisos")
include_router_if_available(api_status, "api_status")
include_router_if_available(api_equipamentos_ti, "api_equipamentos_ti")
include_router_if_available(api_cofre, "api_cofre")
include_router_if_available(api_whatsapp, "api_whatsapp")
include_router_if_available(api_tenants, "api_tenants")

# ─────────────────────────────────────────────────────────────────────────────
# 🤖 Endpoint: Chat do Assistente Contador
# ─────────────────────────────────────────────────────────────────────────────
from fastapi.responses import JSONResponse
from pydantic import BaseModel as PydanticBaseModel

class ChatRequest(PydanticBaseModel):
    mensagem: str

RESPOSTAS_IA = [
    # ---- DÚVIDAS DO PORTAL / TI ----
    (["impressora","imprimir","impressão","toner","papel","offline","online","equipamento"],
     "🖨️ Para verificar o status das impressoras, acesse o menu **Impressoras**. Lá você vê quais estão online, nível de toner e os últimos trabalhos. Se uma impressora estiver offline, tente reiniciá-la ou abra um chamado!"),
    (["chamado","suporte","problema","ticket","abrir","atendimento","solicitação","reclamação"],
     "🎫 Abra um chamado de suporte clicando em **Chamados** no menu lateral. Descreva o problema detalhadamente (ex: 'Impressora RH atolou') para agilizar o atendimento. Nosso time técnico responde rápido!"),
    (["usuário","usuario","acesso","login","senha","perfil","permissão","cadastro","bloqueado"],
     "👤 Questões de usuários e acesso estão em **Usuários** no painel administrativo. Se alguém esqueceu a senha, há a opção de reset na tela de login ou você mesmo pode gerar uma nova pelo painel."),
    (["relatorio","relatório","gráfico","estatística","histórico","dados","exportar","excel"],
     "📊 Os relatórios completos de impressão e custos ficam na seção **Relatórios**. Você pode filtrar por data, departamento ou usuário, e depois exportar tudo para Excel!"),
    (["custo","valor","fatura","financeiro","pagamento","gasto","preço","dinheiro","real"],
     "💰 O controle financeiro e custos de impressão estão na seção **Financeiro**. Lá você acompanha os gastos totais e por departamento (Rateio) de cada período."),
    (["rede","maquina","máquina","computador","pc","monitor","conectado"],
     "🖥️ O **Monitor de Rede** mostra em tempo real quais máquinas e IPs estão conectados. Acesse pelo menu principal para ver o status da rede e tempo de inatividade."),
    (["cofre","senha","credencial","segurança","chave","secret"],
     "🔐 O **Cofre de Senhas** é o local seguro para guardar credenciais corporativas, senhas de roteadores e do certificado digital. Fica criptografado e apenas admins acessam."),
    (["dashboard","painel","inicio","início","resumo","geral"],
     "📈 O **Dashboard** é a tela inicial com os indicadores principais: máquinas online, impressoras ativas, páginas totais e custos. Ele atualiza automaticamente a cada 5 segundos!"),
    (["configuração","configuracoes","configurar","definição","logo","cor","tema"],
     "⚙️ As **Configurações** do sistema ficam no menu lateral (ícone de engrenagem). Lá você altera o nome do escritório, logo, paleta de cores e regras de negócio do Portal."),

    # ---- DÚVIDAS CONTÁBEIS / RH / FISCAL ----
    (["mei","microempreendedor","das","limite mei","desenquadramento"],
     "🏢 **MEI (Microempreendedor Individual)**:\n• O limite de faturamento atual é R$ 81.000,00 anuais (proporcional no ano de abertura).\n• O imposto (DAS) vence todo dia 20 e é um valor fixo mensal.\n• Pode ter no máximo 1 empregado registrado."),
    (["simples nacional","simples","anexo","das simples","alíquota","aliquota"],
     "📄 **Simples Nacional**:\n• Regime tributário simplificado que une federais, estaduais e municipais em 1 guia (DAS).\n• Vence no dia 20 do mês seguinte ao faturamento.\n• Possui 5 anexos (I: Comércio, II: Indústria, III a V: Serviços). A alíquota inicial varia de 4% a 15,5%."),
    (["férias","ferias","1/3","abono pecuniário","vender férias"],
     "🏖️ **Férias**:\n• O empregado tem direito a 30 dias após 12 meses trabalhados (período aquisitivo).\n• O pagamento deve ser feito até 2 dias ANTES do início das férias.\n• É acrescido do terço constitucional (1/3).\n• Pode vender até 10 dias (abono pecuniário)."),
    (["décimo terceiro","decimo terceiro","13","13º"],
     "🎄 **13º Salário**:\n• A 1ª parcela (50% sem descontos) deve ser paga até 30 de novembro.\n• A 2ª parcela (com descontos de INSS e IRRF) deve ser paga até 20 de dezembro.\n• Cada mês com 15 dias ou mais trabalhados dá direito a 1/12 avos."),
    (["rescisão","rescisao","demissão","demissao","aviso prévio","justa causa","seguro desemprego"],
     "🛑 **Rescisão de Contrato**:\n• Em dispensa *sem justa causa*, paga-se: saldo de salário, férias proporcionais + 1/3, 13º proporcional e multa de 40% do FGTS.\n• O aviso prévio pode ser trabalhado (30 a 90 dias) ou indenizado.\n• O pagamento da rescisão deve ser feito em até 10 dias após o término do contrato."),
    (["fgts","fundo de garantia","saque aniversário","multa fgts"],
     "🏗️ **FGTS**:\n• Recolhimento mensal de 8% do salário (2% para Jovem Aprendiz).\n• Vence até o dia 20 do mês seguinte (pelo eSocial/FGTS Digital).\n• Em caso de dispensa sem justa causa, a empresa paga 40% de multa sobre o saldo."),
    (["inss","previdência","aposentadoria","contribuição","desconto inss"],
     "👴 **INSS**:\n• As alíquotas de desconto para empregados são progressivas: 7,5%, 9%, 12% ou 14% de acordo com a faixa salarial.\n• O limite máximo de desconto segue o teto do INSS do ano vigente.\n• O vencimento da guia (DARF Previdenciário) é até o dia 20."),
    (["irpf","imposto de renda","malha fina","restituição","carnê leão"],
     "🦁 **Imposto de Renda Pessoa Física (IRPF)**:\n• O prazo de entrega anual geralmente vai de março a maio.\n• Fica obrigado a declarar quem obteve rendimentos tributáveis acima do limite isento no ano base.\n• O IRRF (retido na fonte) começa a descontar nos salários acima de R$ 2.824,00 (tabela de 2024), descontando a parcela a deduzir."),
    (["certidão","certidao","cnd","negativa","regularidade"],
     "📜 **Certidões (CND)**:\n• A Certidão Negativa de Débitos atesta que a empresa não tem dívidas.\n• **Receita Federal**: Dívidas tributárias e previdenciárias.\n• **FGTS**: Emitida pela Caixa Econômica.\n• **Trabalhista (CNDT)**: Emitida pelo TST."),
    (["esocial","e-social","folha","fechamento","dctfweb"],
     "💻 **eSocial & DCTFWeb**:\n• O eSocial unificou a entrega de informações trabalhistas e previdenciárias.\n• A folha de pagamento deve ser fechada e transmitida até o dia 15.\n• Os tributos (INSS) são gerados pelo DCTFWeb com vencimento para o dia 20."),

    # ---- BÁSICOS ----
    (["oi","olá","ola","bom dia","boa tarde","boa noite","tudo bem","tudo bom","hey","hello"],
     "😊 Olá! Tudo ótimo por aqui! Sou o Contadorzinho, seu assistente virtual de TI e Contabilidade. Pode me perguntar sobre impressoras, relatórios, MEI, rescisões, Simples Nacional... o que precisar!"),
    (["quem","você","voce","seu nome","nome","assistente","ajuda","o que faz","chatbot","ia","bot"],
     "🧑‍💼 Sou o **Contadorzinho**, o assistente oficial do Portal de Acesso! Fui treinado com dezenas de regras contábeis, trabalhistas e de TI para tirar suas dúvidas rápidas do dia a dia do escritório. Do Simples Nacional à impressora sem papel, é só mandar a dúvida!"),
    (["obrigado","obrigada","valeu","thanks","grato","grata","ajudou","top","perfeito"],
     "😄 Eu que agradeço! Se surgir mais alguma dúvida no escritório, é só me chamar. Bom trabalho! 👋"),
]

def responder_contador(mensagem: str) -> str:
    deepseek_key = os.getenv("DEEPSEEK_API_KEY", "").strip()
    if deepseek_key:
        try:
            from openai import OpenAI
            client = OpenAI(api_key=deepseek_key, base_url="https://api.deepseek.com")
            prompt = (
                "Você é o 'Contadorzinho', um assistente virtual super inteligente, experiente e simpático "
                "de um escritório de contabilidade (Portal de Acesso). "
                "Sua função é responder dúvidas sobre contabilidade (tributária, trabalhista, societária, MEI, Simples Nacional, etc), "
                "bem como funções de TI (impressoras, chamados, rede). "
                "Seja direto, prestativo e use alguns emojis para deixar o texto leve. "
                "Sempre formate em Markdown com listas e negrito quando facilitar a leitura."
            )
            response = client.chat.completions.create(
                model="deepseek-chat",
                messages=[
                    {"role": "system", "content": prompt},
                    {"role": "user", "content": mensagem}
                ],
                max_tokens=600,
                temperature=0.3
            )
            return response.choices[0].message.content
        except Exception as e:
            return f"❌ Ops, erro na conexão com o DeepSeek: {str(e)}"

    # Fallback local se não houver chave do DeepSeek
    texto = mensagem.lower().strip()
    for palavras, resposta in RESPOSTAS_IA:
        if any(p in texto for p in palavras):
            return resposta
    return ("🤔 Minha conexão com o **DeepSeek** ainda não está ativa (falta a variável `DEEPSEEK_API_KEY` no `.env`). "
            "Por enquanto só sei responder sobre: **impressoras**, **chamados**, **relatórios** ou **usuários**!")

@app.post("/api/contador-chat")
async def contador_chat(req: ChatRequest, request: Request):
    if not req.mensagem or len(req.mensagem.strip()) < 1:
        return JSONResponse({"ok": False, "error": "Mensagem vazia"})
    resposta = responder_contador(req.mensagem)
    return JSONResponse({"ok": True, "resposta": resposta})

# Montar Arquivos EstÃ¡ticos (CSS, JS, etc.)
app.mount("/static", StaticFiles(directory=os.path.join(BASE_DIR, "static")), name="static")
# Montar Pasta de Uploads (Logo, Documentos, etc.)
app.mount("/uploads", StaticFiles(directory=os.path.join(BASE_DIR, "public_html", "uploads")), name="uploads")

# Inicializar Banco
@app.on_event("startup")
def on_startup():
    import time
    last_error = None
    for attempt in range(30):
        try:
            SQLModel.metadata.create_all(engine)
            ensure_schema_compatibility()
            print("[STARTUP] Banco inicializado com sucesso.")
            return
        except Exception as e:
            last_error = e
            print(f"[STARTUP] Banco indisponÃ­vel (tentativa {attempt + 1}/30): {e}")
            time.sleep(2)

    print(f"[STARTUP] Seguindo sem criar tabelas apÃ³s falha persistente: {last_error}")



AGENT_API_KEY = os.getenv("AGENT_API_KEY")

@app.get("/api/agent/download_latest")
async def download_latest_agent(x_api_key: Optional[str] = Header(None)):
    """Entrega o executável atual para o auto-update do agente."""
    if not AGENT_API_KEY or x_api_key != AGENT_API_KEY:
        raise HTTPException(status_code=401, detail="Não autorizado")

    candidates = [
        os.path.abspath(os.path.join(BASE_DIR, "..", "dist", "agent_2.0.exe")),
        os.path.abspath(os.path.join(BASE_DIR, "..", "agent_2.0.exe")),
    ]
    exe_path = next((path for path in candidates if os.path.isfile(path)), None)
    if not exe_path:
        raise HTTPException(status_code=404, detail="Executável do agente não encontrado")

    return FileResponse(
        exe_path,
        media_type="application/vnd.microsoft.portable-executable",
        filename="agent_2.0.exe",
    )

def update_status_files(maquina: str, ip: str, payload: dict, versao: str):
    """Atualiza os arquivos JSON de status exportando do DB para manter compatibilidade"""
    from models import StatusMaquina
    agora = datetime.now(timezone.utc)
    status_file = "status_maquinas.json"
    
    with Session(engine) as session:
        # Atualizar no DB (Single Source of Truth)
        stat = session.get(StatusMaquina, ip)
        if not stat:
            stat = StatusMaquina(ip=ip)
        stat.nome = maquina
        stat.status = "Online"
        stat.windows_user = payload.get("windows_user", "Desconhecido")
        stat.timestamp = agora
        stat.versao_agente = versao
        session.add(stat)
        
        # Remove registros antigos da mesma máquina com IP diferente (evitar duplicação Offline/Online)
        old_stats = session.exec(select(StatusMaquina).where(StatusMaquina.nome == maquina).where(StatusMaquina.ip != ip)).all()
        for old in old_stats:
            session.delete(old)
            
        session.commit()
        
        # Gerar o JSON completo a partir do DB apenas a cada 5 segundos
        # Para evitar que mÃºltiplos workers destruam o disco, usaremos probabilidade simples 
        # ou apenas tentamos gravar e ignoramos erros. O DB Ã© a fonte confiÃ¡vel.
        import time
        global last_flush_time
        if 'last_flush_time' not in globals():
            last_flush_time = 0
            
        current_time = time.time()
        if current_time - last_flush_time > 5:
            last_flush_time = current_time
            try:
                all_stats = session.exec(select(StatusMaquina)).all()
                data = {"Maquinas": {}}
                for s in all_stats:
                    data["Maquinas"][s.ip] = {
                        "Nome": s.nome,
                        "Status": "Online" if (agora - s.timestamp).total_seconds() < 300 else "Offline",
                        "WindowsUser": s.windows_user,
                        "Timestamp": s.timestamp.strftime("%Y-%m-%d %H:%M:%S"),
                        "VersaoAgente": s.versao_agente
                    }
                data["UltimaAtualizacao"] = agora.strftime("%Y-%m-%d %H:%M:%S")
                with open(status_file, "w") as f:
                    json.dump(data, f, indent=4)
            except: pass

def classify_doc_type(doc_name: str, app_process: str) -> str:
    """Classifica o tipo de documento (mesma lÃ³gica do PHP)"""
    doc_lower = doc_name.lower()
    app_lower = app_process.lower()
    
    if any(x in app_lower for x in ["winword", "word"]): return "Word"
    if any(x in app_lower for x in ["excel", "xls"]): return "Excel"
    if any(x in app_lower for x in ["powerpnt", "powerpoint"]): return "PPT"
    if any(x in app_lower for x in ["acrord", "pdf"]): return "PDF"
    if any(x in app_lower for x in ["chrome", "msedge", "firefox"]): return "Web"
    
    # Fallback por extensÃ£o/nome
    if ".pdf" in doc_lower: return "PDF"
    if ".doc" in doc_lower or ".docx" in doc_lower: return "Word"
    if ".xls" in doc_lower or ".xlsx" in doc_lower: return "Excel"
    
    return "Sistema"

@app.post("/api/v1/agent")
@limiter.limit("60/minute")
async def agent_endpoint(
    request: Request,
    background_tasks: BackgroundTasks,
    x_api_key: Optional[str] = Header(None)
):
    if not AGENT_API_KEY or x_api_key != AGENT_API_KEY:
        raise HTTPException(status_code=401, detail="NÃ£o autorizado")
    
    data = await request.json()
    maquina = data.get("maquina")
    ip = data.get("ip")
    versao = data.get("versao", "2.0")
    payload = data.get("dados", {})
    tipo = payload.get("tipo")
    
    if tipo == "heartbeat":
        background_tasks.add_task(update_status_files, maquina, ip, payload, versao)
        return {
            "success": True, 
            "latest_version": "2.4.3",
            "commands": []
        }
    
    if tipo == "impressoes":
        from notifications import send_alert
        with Session(engine) as session:
            for job in payload.get("jobs", []):
                # Busca impressora
                printer_name = job.get("PrinterName", "Padrao")
                statement = select(Impressora).where(Impressora.nome == printer_name)
                printer = session.exec(statement).first()
                
                if not printer:
                    printer = Impressora(nome=printer_name, ip=ip, status="Online")
                    session.add(printer)
                    session.commit()
                    session.refresh(printer)
                
                # Classifica e Salva
                doc_name = job.get("DocumentName", "Documento")
                app_proc = job.get("AppProcess", "")
                tipo_doc = classify_doc_type(doc_name, app_proc)
                paginas = int(job.get("TotalPages", 1))
                digital = bool(job.get("Digital", False))
                preco_unitario = 0.0 if digital else 0.10
                
                nova_imp = Impressao(
                    usuario=job.get("UserName", "Sistema"),
                    impressora_id=printer.id,
                    maquina=maquina,
                    documento=doc_name,
                    tipo_documento=tipo_doc,
                    paginas=paginas,
                    tipo_impressao="Digital" if digital else "Preto e Branco",
                    ip_maquina=ip,
                    preco_unitario=preco_unitario,
                    valor_total=paginas * preco_unitario
                )
                session.add(nova_imp)
                
                # Alerta de volume alto (ex: 50+ paginas)
                if paginas >= 50:
                    msg = (
                        f"⚠️ *ALERTA DE ALTO VOLUME*\n\n"
                        f"👤 Usuário: {nova_imp.usuario}\n"
                        f"🖨️ Impressora: {printer.nome}\n"
                        f"📝 Documento: {doc_name}\n"
                        f"📄 Páginas: {paginas}"
                    )
                    background_tasks.add_task(send_alert, msg)

            session.commit()
            
    return {"success": True}

# --- ROTA DE COMPATIBILIDADE: Agentes Antigos (v1/v2) ---
# Agentes instalados com API_URL="http://.../api_agente_v2.php" ainda chamam esta rota.
# Redirecionamos internamente para o mesmo handler sem exigir atualizaÃ§Ã£o do agente.
@app.post("/api_agente_v2.php")
@limiter.limit("60/minute")
async def agent_compat_endpoint(
    request: Request,
    background_tasks: BackgroundTasks,
    x_api_key: Optional[str] = Header(None)
):
    """Alias de compatibilidade para agentes legados."""
    return await agent_endpoint(request, background_tasks, x_api_key)


@app.get("/", response_class=HTMLResponse)
async def dashboard(request: Request):
    user_id = request.cookies.get("user_id")
    if not user_id: return RedirectResponse(url="/login")
    if request.cookies.get("user_role") == "estudante":
        return RedirectResponse(url="/estudos")
    
    with Session(engine) as session:
        # Busca o usuÃ¡rio no banco para ter as permissÃµes REAIS e ATUAIS
        user = session.get(Usuario, int(user_id))
        if not user: return RedirectResponse(url="/login")

        # Redirecionamento forÃ§ado em tempo real
        try:
            user_perms = normalize_permissions(json.loads(user.permissoes or "[]"))
            print(f"[FORCE_REDIRECT_DEBUG] User: {user.username}, Role: {user.role}, Perms: {user_perms}")
            
            if user.role == "operator":
                # Se for operador, vamos ver se ele tem Dashboard. Se NÃƒO tiver, manda pra /meu_painel.
                if "Dashboard" not in user_perms:
                    print(f"[FORCE_REDIRECT_ACTION] Redirecting {user.username} to /meu_painel")
                    return RedirectResponse(url="/meu_painel")
        except Exception as e:
            print(f"[FORCE_REDIRECT_ERROR] {str(e)}")
            pass

        hoje = datetime.now().date()
        
        # EstatÃ­sticas do Dashboard Executivo
        total_paginas = session.exec(select(func.sum(Impressao.paginas))).one() or 0
        faturamento = session.exec(select(func.sum(Impressao.valor_total))).one() or 0
        paginas_hoje = session.exec(select(func.sum(Impressao.paginas)).where(func.date(Impressao.data_hora) == hoje)).one() or 0
        
        # MÃ¡quinas na Rede (Dados do Agente)
        from models import StatusMaquina
        try:
            total_maquinas = session.exec(select(func.count(StatusMaquina.ip))).one() or 0
            
            from datetime import timedelta
            limite_online = datetime.now(timezone.utc) - timedelta(minutes=5)
            maquinas_online = session.exec(select(func.count(StatusMaquina.ip)).where(StatusMaquina.timestamp >= limite_online)).one() or 0
        except Exception as e:
            print("Erro ao carregar maquinas:", str(e))
            maquinas_online = 0
            total_maquinas = 0

        # EstatÃ­sticas de Chamados (TI)
        from models import Chamado
        total_chamados = session.exec(select(func.count(Chamado.id))).one() or 0
        urgentes_chamados = session.exec(
            select(func.count(Chamado.id)).where(
                Chamado.status == 'Aberto',
                Chamado.data_abertura < (datetime.now(timezone.utc) - timedelta(hours=4))
            )
        ).one() or 0
        stats_chamados = {
            "abertos": session.exec(select(func.count(Chamado.id)).where(Chamado.status == 'Aberto')).one(),
            "atendimento": session.exec(select(func.count(Chamado.id)).where(Chamado.status == 'Em Atendimento')).one(),
            "resolvidos": session.exec(select(func.count(Chamado.id)).where(Chamado.status == 'Resolvido')).one(),
            "total": total_chamados,
            "urgentes": urgentes_chamados,
            "saude": round((session.exec(select(func.count(Chamado.id)).where(Chamado.status == 'Resolvido')).one() or 0) / total_chamados * 100) if total_chamados > 0 else 100
        }

        # Dados para GrÃ¡ficos
        # Top 5 Impressoras
        stmt_imp = select(Impressora.nome, func.sum(Impressao.paginas).label("total")).join(Impressao, Impressao.impressora_id == Impressora.id).group_by(Impressora.nome).order_by(desc(func.sum(Impressao.paginas))).limit(5)
        top_impressoras = session.exec(stmt_imp).all()
        
        # Top 5 UsuÃ¡rios
        stmt_usu = select(Impressao.usuario, func.sum(Impressao.paginas).label("total")).group_by(Impressao.usuario).order_by(desc(func.sum(Impressao.paginas))).limit(5)
        top_usuarios = session.exec(stmt_usu).all()

        # Atividade Recente Detalhada
        atividades_raw = session.exec(select(Impressao).order_by(desc(Impressao.data_hora)).limit(10)).all()
        atividades_fmt = []
        for imp in atividades_raw:
            # ProteÃ§Ã£o para valores do PHP (ex: "1,50" em vez de 1.50)
            try:
                v = float(str(imp.valor_total).replace(",", ".")) if imp.valor_total else 0.0
            except:
                v = 0.0
            
            # ProteÃ§Ã£o para datas que podem vir como string do SQLite
            try:
                if isinstance(imp.data_hora, str):
                    dt = datetime.strptime(imp.data_hora[:19].replace("T", " "), "%Y-%m-%d %H:%M:%S")
                else:
                    dt = imp.data_hora
                d_str = dt.strftime('%d/%m')
                h_str = dt.strftime('%H:%M')
            except:
                d_str = "--/--"
                h_str = "--:--"
                
            atividades_fmt.append({
                "data_str": d_str,
                "hora_str": h_str,
                "tipo_documento": imp.tipo_documento,
                "tipo_impressao": imp.tipo_impressao,
                "documento": imp.documento,
                "maquina": imp.maquina,
                "usuario": imp.usuario,
                "paginas": imp.paginas,
                "tipo_impressao": imp.tipo_impressao or "Preto e Branco",
                "valor_fmt": f"{v:.2f}".replace(".", ",")
            })

        return templates.TemplateResponse(
            request=request,
            name="dashboard.html",
            context={
                "config": APP_CONFIG,
                "user_perms_real": user_perms,
                "stats": {
                    "paginas_total": total_paginas,
                    "faturamento": faturamento,
                    "paginas_hoje": paginas_hoje,
                    "maquinas_online": maquinas_online,
                    "total_maquinas": total_maquinas,
                    "impressoras_online": session.exec(select(func.count(Impressora.id)).where(Impressora.status == "Online")).one(),
                    "chamados": stats_chamados
                },
                "top_impressoras": [{"nome": i[0] or "Local", "total": int(i[1] or 0)} for i in top_impressoras],
                "top_usuarios": [{"nome": u[0] or "Desconhecido", "total": int(u[1] or 0)} for u in top_usuarios],
                "atividades": atividades_fmt,
                "hoje": hoje.strftime('%d/%m/%Y')
            }
        )

# --- APIs COMPATÃVEIS COM O DASHBOARD ---

@app.get("/api/stats_impressao")
async def api_stats_impressao():
    with Session(engine) as session:
        # Top 5 Impressoras
        # Nota: JOIN com Impressora para pegar o nome
        results = session.exec(
            select(Impressora.nome, func.sum(Impressao.paginas).label("total"))
            .join(Impressao, Impressao.impressora_id == Impressora.id)
            .group_by(Impressora.nome)
            .order_by(func.sum(Impressao.paginas).desc())
            .limit(5)
        ).all()
        por_impressora = [{"nome": r[0], "total": int(r[1])} for r in results]

        # Top 5 UsuÃ¡rios
        results_u = session.exec(
            select(Impressao.usuario, func.sum(Impressao.paginas))
            .group_by(Impressao.usuario)
            .order_by(func.sum(Impressao.paginas).desc())
            .limit(5)
        ).all()
        por_usuario = [{"nome": r[0], "total": int(r[1])} for r in results_u]

        return {
            "success": True,
            "stats": {
                "impressoras": por_impressora,
                "usuarios": por_usuario
            }
        }

@app.get("/api/recent_activity")
async def api_recent_activity():
    with Session(engine) as session:
        # JOIN com Impressora para pegar o nome_impressora
        statement = select(Impressao, Impressora.nome).join(Impressora, isouter=True).order_by(Impressao.data_hora.desc()).limit(8)
        results = session.exec(statement).all()
        
        atividades = []
        for imp, nome_imp in results:
            atividades.append({
                "id": imp.id,
                "usuario": imp.usuario,
                "maquina": imp.maquina,
                "documento": imp.documento,
                "tipo_documento": imp.tipo_documento,
                "paginas": imp.paginas,
                "tipo_impressao": imp.tipo_impressao,
                "valor_total": imp.valor_total,
                "nome_impressora": nome_imp or "Local",
                "data_formatada": imp.data_hora.strftime("%d/%m"),
                "hora_formatada": imp.data_hora.strftime("%H:%M:%S")
            })

        return {
            "success": True,
            "atividades": atividades
        }

# --- FINANCEIRO ---

@app.get("/financeiro", response_class=HTMLResponse)
def view_financeiro(request: Request):
    user_id = get_signed_cookie(request, "user_id")
    user_role = get_signed_cookie(request, "user_role")
    if not user_id:
        return RedirectResponse(url="/login")
    if user_role not in ["admin", "gestor"]:
        return RedirectResponse(url="/")
    return templates.TemplateResponse(request, "financeiro.html", {"config": APP_CONFIG})

# --- ROTAS DE LICENCIAMENTO ---

@app.get("/ativar", response_class=HTMLResponse)
async def view_ativar(request: Request, error: str = None):
    user_id = request.cookies.get("user_id")
    user_role = request.cookies.get("user_role")
    if not user_id: return RedirectResponse(url="/login")
    if user_role != "admin": return RedirectResponse(url="/")
    return templates.TemplateResponse(
        request=request,
        name="ativar.html",
        context={
            "hardware_id": license_service.get_hardware_id(),
            "error": error
        }
    )

@app.post("/ativar")
async def process_ativar(request: Request, token: str = Form(...)):
    user_id = request.cookies.get("user_id")
    user_role = request.cookies.get("user_role")
    if not user_id: return RedirectResponse(url="/login", status_code=303)
    if user_role != "admin": return RedirectResponse(url="/", status_code=303)
    if license_service.activate_token(token):
        return RedirectResponse(url="/", status_code=303)
    return RedirectResponse(url="/ativar?error=Token InvÃ¡lido ou jÃ¡ utilizado", status_code=303)

# --- OUTRAS TELAS MIGRADAS ---

@app.get("/impressoras", response_class=HTMLResponse)
async def view_impressoras(request: Request):
    user_id = request.cookies.get("user_id")
    user_role = request.cookies.get("user_role")
    user_perms = request.cookies.get("user_perms") or "[]"
    if not user_id: return RedirectResponse(url="/login")
    try:
        perms = json.loads(user_perms)
    except:
        perms = []
    if user_role != "admin" and "Equipamentos" not in perms:
        return RedirectResponse(url="/")
    
    with Session(engine) as session:
        lista = session.exec(select(Impressora)).all()
        return templates.TemplateResponse(
            request=request,
            name="impressoras.html",
            context={
                "config": APP_CONFIG,
                "impressoras": lista
            }
        )

@app.get("/relatorios", response_class=HTMLResponse)
async def view_relatorios(request: Request, mes: Optional[str] = None, ano: Optional[str] = None, usuario: Optional[str] = None, impressora_id: Optional[str] = None, tipo: Optional[str] = None):
    user_id = request.cookies.get("user_id")
    user_role = request.cookies.get("user_role")
    user_perms = request.cookies.get("user_perms") or "[]"
    if not user_id: return RedirectResponse(url="/login")
    try:
        perms = json.loads(user_perms)
    except:
        perms = []
    if user_role != "admin" and "RelatÃ³rios" not in perms:
        return RedirectResponse(url="/")
    
    agora = datetime.now()
    mes = mes or agora.strftime("%m")
    ano = ano or agora.strftime("%Y")

    impressora_id_value = None
    if impressora_id:
        try:
            impressora_id_value = int(impressora_id)
        except (TypeError, ValueError):
            impressora_id_value = None
    
    with Session(engine) as session:
        # Filtro de data para o mÃªs selecionado
        start_date = datetime(int(ano), int(mes), 1)
        if int(mes) == 12:
            end_date = datetime(int(ano) + 1, 1, 1)
        else:
            end_date = datetime(int(ano), int(mes) + 1, 1)

        filtros = [Impressao.data_hora >= start_date, Impressao.data_hora < end_date]
        if usuario:
            filtros.append(Impressao.usuario == usuario)
        if impressora_id_value:
            filtros.append(Impressao.impressora_id == impressora_id_value)
        if tipo in ("Digital", "Fisica"):
            if tipo == "Digital":
                filtros.append(Impressao.tipo_impressao == "Digital")
            else:
                filtros.append(Impressao.tipo_impressao.in_(["Preto e Branco", "Colorida", "P&B", "Colorido"]))
        impressoras_filtro = session.exec(select(Impressora).order_by(Impressora.nome)).all()
        usuarios_filtro = session.exec(select(Impressao.usuario).distinct().order_by(Impressao.usuario)).all()

        # Resumo por Impressora
        stmt_imp = select(Impressora.nome, func.sum(Impressao.paginas), func.sum(Impressao.valor_total)).join(Impressao, Impressao.impressora_id == Impressora.id).where(*filtros).group_by(Impressora.nome).order_by(desc(func.sum(Impressao.paginas)))
        resumo_impressoras = session.exec(stmt_imp).all()
        
        # Resumo por UsuÃ¡rio
        stmt_usu = select(Impressao.usuario, func.sum(Impressao.paginas), func.sum(Impressao.valor_total)).where(*filtros).group_by(Impressao.usuario).order_by(desc(func.sum(Impressao.paginas)))
        resumo_usuarios = session.exec(stmt_usu).all()

        digital_pages = session.exec(
            select(func.sum(Impressao.paginas))
            .where(*filtros, Impressao.tipo_impressao == "Digital")
        ).one() or 0
        digital_jobs = session.exec(
            select(func.count(Impressao.id))
            .where(Impressao.tipo_impressao == "Digital")
            .where(*filtros)
        ).one() or 0
        physical_pages = session.exec(
            select(func.sum(Impressao.paginas))
            .where(*filtros, Impressao.tipo_impressao != "Digital")
        ).one() or 0
        physical_jobs = session.exec(
            select(func.count(Impressao.id))
            .where(Impressao.tipo_impressao != "Digital")
            .where(*filtros)
        ).one() or 0

        # EvoluÃ§Ã£o DiÃ¡ria (para o grÃ¡fico)
        stmt_evolu = select(func.date(Impressao.data_hora), func.sum(Impressao.paginas)).where(*filtros).group_by(func.date(Impressao.data_hora)).order_by(func.date(Impressao.data_hora))
        evolucao_diaria = session.exec(stmt_evolu).all()
        
        evolucao_fmt = []
        for e in evolucao_diaria:
            try:
                # e[0] Ã© uma string do SQLite date(), ex: "2026-05-11"
                d_obj = datetime.strptime(str(e[0]), "%Y-%m-%d")
                evolucao_fmt.append({"data": d_obj.strftime("%d/%m"), "paginas": int(e[1] or 0)})
            except:
                evolucao_fmt.append({"data": str(e[0]), "paginas": int(e[1] or 0)})

        # Auditoria Detalhada (200 registros)
        atividades_raw = session.exec(select(Impressao).where(*filtros).order_by(desc(Impressao.data_hora)).limit(200)).all()
        atividades_fmt = []
        for imp in atividades_raw:
            try:
                v = float(str(imp.valor_total).replace(",", ".")) if imp.valor_total else 0.0
            except:
                v = 0.0
                
            try:
                if isinstance(imp.data_hora, str):
                    dt = datetime.strptime(imp.data_hora[:19].replace("T", " "), "%Y-%m-%d %H:%M:%S")
                else:
                    dt = imp.data_hora
                d_str = dt.strftime('%d/%m/%Y %H:%M')
            except:
                d_str = "--/--"
                
            atividades_fmt.append({
                "data_hora_str": d_str,
                "tipo_documento": imp.tipo_documento,
                "documento": imp.documento,
                "usuario": imp.usuario,
                "paginas": imp.paginas,
                "valor_fmt": f"{v:.2f}".replace(".", ",")
            })

        equipamentos_ti = session.exec(
            select(EquipamentoTI).order_by(EquipamentoTI.colaborador, EquipamentoTI.id.desc())
        ).all()
        equipamentos_fmt = []
        for eq in equipamentos_ti:
            nome_computador = str(getattr(eq, "nome_computador", "") or "").strip()
            equipamentos_fmt.append({
                "id": eq.id,
                "colaborador": eq.colaborador,
                "nome_computador": nome_computador,
                "nome_computador_bruto": nome_computador,
                "nome_computador_display": nome_computador or "Não informado",
                "setor": eq.setor,
                "tipo_equipamento": eq.tipo_equipamento,
                "patrimonio": eq.patrimonio,
                "data_registro": eq.data_registro.strftime("%d/%m/%Y %H:%M") if eq.data_registro else "",
            })
        equipamentos_sem_nome = sum(
            1 for eq in equipamentos_ti if not str(getattr(eq, "nome_computador", "") or "").strip()
        )

        equipamentos_por_tipo = {}
        contagem_por_tipo = {}
        for eq in equipamentos_fmt:
            tipo_eq = (eq["tipo_equipamento"] or "Outros").strip()
            if tipo_eq not in equipamentos_por_tipo:
                equipamentos_por_tipo[tipo_eq] = []
                contagem_por_tipo[tipo_eq] = 0
            equipamentos_por_tipo[tipo_eq].append(eq)
            contagem_por_tipo[tipo_eq] += 1
        
        return templates.TemplateResponse(
            request=request,
            name="relatorios.html",
            context={
                "config": APP_CONFIG,
                "mes_atual": mes,
                "ano_atual": ano,
                "filtro_usuario": usuario or "",
                "filtro_impressora_id": impressora_id_value or "",
                "filtro_tipo": tipo or "",
                "usuarios_filtro": usuarios_filtro,
                "impressoras_filtro": impressoras_filtro,
                "resumo_impressoras": [{"nome": i[0] or "Local", "paginas": int(i[1] or 0), "valor": float(str(i[2] or 0).replace(",", "."))} for i in resumo_impressoras],
                "resumo_usuarios": [{"nome": u[0] or "Desconhecido", "paginas": int(u[1] or 0), "valor": float(str(u[2] or 0).replace(",", "."))} for u in resumo_usuarios],
                "resumo_fisico": {"trabalhos": int(physical_jobs), "paginas": int(physical_pages)},
                "resumo_digital": {"trabalhos": int(digital_jobs), "paginas": int(digital_pages), "valor": 0.0},
                "evolucao": evolucao_fmt,
                "atividades": atividades_fmt,
                "equipamentos_ti": equipamentos_fmt,
                "equipamentos_por_tipo": equipamentos_por_tipo,
                "contagem_por_tipo": contagem_por_tipo,
                "total_equipamentos_ti": len(equipamentos_ti),
                "equipamentos_sem_nome": equipamentos_sem_nome,
            }
        )

@app.get("/monitor_usuarios", response_class=HTMLResponse)
async def view_monitor_usuarios(request: Request):
    user_id = request.cookies.get("user_id")
    user_role = request.cookies.get("user_role")
    user_perms = request.cookies.get("user_perms") or "[]"
    if not user_id: return RedirectResponse(url="/login")
    try:
        perms = json.loads(user_perms)
    except:
        perms = []
    if user_role != "admin" and "UsuÃ¡rios" not in perms:
        return RedirectResponse(url="/")
    return templates.TemplateResponse(
        request=request,
        name="monitor_usuarios.html",
        context={"config": APP_CONFIG}
    )

@app.get("/equipamentos_ti", response_class=HTMLResponse)
async def view_equipamentos_ti(request: Request):
    try:
        user_id = request.cookies.get("user_id")
        user_role = request.cookies.get("user_role")

        if not user_id:
            return RedirectResponse(url="/login")
        if user_role != "admin":
            return RedirectResponse(url="/")

        equipamentos = []
        equipamentos_sem_nome = 0
        equipamentos_fmt = []
        equipamentos_load_error = None

        try:
            with Session(engine) as session:
                equipamentos = session.exec(select(EquipamentoTI).order_by(EquipamentoTI.colaborador)).all()
                equipamentos_sem_nome = sum(
                    1 for eq in equipamentos if not str(getattr(eq, "nome_computador", "") or "").strip()
                )
                for eq in equipamentos:
                    nome_computador = str(getattr(eq, "nome_computador", "") or "").strip()
                    equipamentos_fmt.append({
                        "id": eq.id,
                        "colaborador": eq.colaborador,
                        "nome_computador": nome_computador,
                        "nome_computador_bruto": nome_computador,
                        "nome_computador_display": nome_computador or "Não informado",
                        "setor": eq.setor,
                        "tipo_equipamento": eq.tipo_equipamento,
                        "marca": eq.marca,
                        "modelo": eq.modelo,
                        "modelo_monitor": eq.modelo_monitor,
                        "quantidade_monitores": eq.quantidade_monitores,
                        "processador": eq.processador,
                        "memoria_ram": eq.memoria_ram,
                        "sistema_operacional": eq.sistema_operacional,
                        "armazenamento": eq.armazenamento,
                        "patrimonio": eq.patrimonio,
                        "observacoes": eq.observacoes,
                        "data_registro": eq.data_registro.strftime("%d/%m/%Y %H:%M") if eq.data_registro else "",
                    })
        except Exception as db_exc:
            print(f"[equipamentos_ti] Falha ao carregar equipamentos: {db_exc}")
            print(traceback.format_exc())
            equipamentos_load_error = "Nao foi possivel carregar os dados agora. A tela foi aberta sem os registros."

        render_context = {
            "request": request,
            "config": APP_CONFIG,
            "is_admin": True,
            "equipamentos": jsonable_encoder(equipamentos_fmt),
            "total_equipamentos_ti": len(equipamentos_fmt),
            "equipamentos_sem_nome": equipamentos_sem_nome,
            "equipamentos_load_error": equipamentos_load_error,
        }
        html = templates.get_template("equipamentos_ti.html").render(render_context)
        return HTMLResponse(content=html)
    except Exception:
        import traceback
        err = traceback.format_exc()
        print(f"CRITICAL ERROR IN VIEW_EQUIPAMENTOS_TI:\n{err}")
        return HTMLResponse(content=f"<h1>DEBUG ERROR</h1><pre>{err}</pre>", status_code=500)
@app.get("/marmitas", response_class=HTMLResponse)
async def view_marmitas(request: Request):
    return RedirectResponse(url="/")

@app.get("/usuarios", response_class=HTMLResponse)
async def view_usuarios(request: Request):
    user_id = request.cookies.get("user_id")
    user_role = request.cookies.get("user_role")
    user_perms = request.cookies.get("user_perms") or "[]"
    if not user_id: return RedirectResponse(url="/login")
    try:
        perms = json.loads(user_perms)
    except:
        perms = []
    if user_role != "admin" and "UsuÃ¡rios" not in perms:
        return RedirectResponse(url="/")
    with Session(engine) as session:
        users = session.exec(select(Usuario).order_by(Usuario.created_at.desc())).all()
        lista_usuarios = []
        for u in users:
            d = u.dict()
            d["created_at"] = d["created_at"].strftime("%Y-%m-%d %H:%M:%S") if d["created_at"] else None
            # Pre-processa as permissÃµes para evitar erro no template
            try:
                d["permissoes_list"] = json.loads(d.get("permissoes", "[]"))
            except:
                d["permissoes_list"] = []
            lista_usuarios.append(d)
            
        return templates.TemplateResponse(
            request=request,
            name="usuarios.html",
            context={
                "config": APP_CONFIG,
                "usuarios": lista_usuarios
            }
        )

RAIO_X_BANCAS = {
    "Cebraspe": {
        "descricao": "Questões frequentemente julgadas como certo/errado, com atenção a detalhes, exceções e interpretação literal.",
        "foco": ["literalidade da lei", "exceções e prazos", "interpretação de textos", "itens com uma única afirmação"],
        "alertas": ["palavras absolutas", "troca de conceitos próximos", "jurisprudência desatualizada"],
        "fontes": ["https://www.cebraspe.org.br/concursos/", "https://www.gov.br/planalto/pt-br/"],
    },
    "FGV": {
        "descricao": "Costuma explorar interpretação, aplicação prática, alternativas próximas e enunciados contextualizados.",
        "foco": ["interpretação de texto", "casos práticos", "conceitos e exceções", "raciocínio aplicado"],
        "alertas": ["alternativas parcialmente corretas", "comandos negativos", "mistura de regra e exceção"],
        "fontes": ["https://conhecimento.fgv.br/concursos", "https://www.fgv.br/"],
    },
    "FCC": {
        "descricao": "Banca tradicional, com cobrança objetiva de legislação, gramática, conceitos e aplicação direta do conteúdo.",
        "foco": ["lei seca", "conceitos fundamentais", "gramática e interpretação", "jurisprudência consolidada"],
        "alertas": ["artigos parecidos", "prazos e competências", "detalhes de classificação"],
        "fontes": ["https://www.concursosfcc.com.br/", "https://www.planalto.gov.br/"],
    },
    "Vunesp": {
        "descricao": "Perfil geralmente objetivo, com textos, legislação e situações práticas ligadas ao cargo.",
        "foco": ["interpretação", "legislação do edital", "conhecimentos específicos", "situações do cotidiano"],
        "alertas": ["detalhes do texto-base", "conceitos semelhantes", "comando exato da questão"],
        "fontes": ["https://www.vunesp.com.br/", "https://www.planalto.gov.br/"],
    },
    "Instituto AOCP": {
        "descricao": "Provas objetivas com cobrança de conteúdo do edital, legislação e conhecimentos específicos do cargo.",
        "foco": ["legislação", "conhecimentos específicos", "língua portuguesa", "informática e atualidades quando previstas"],
        "alertas": ["artigos e definições", "itens fora do núcleo mais estudado", "confusão entre atribuições"],
        "fontes": ["https://www.institutoaocp.org.br/", "https://www.planalto.gov.br/"],
    },
    "IBFC": {
        "descricao": "Cobrança distribuída pelo edital, combinando teoria, legislação e questões de aplicação.",
        "foco": ["conteúdo programático", "legislação", "português", "conhecimentos específicos"],
        "alertas": ["detalhes de normas", "conceitos próximos", "interpretação do enunciado"],
        "fontes": ["https://concursos.ibfc.org.br/", "https://www.planalto.gov.br/"],
    },
    "Instituto Consulplan": {
        "descricao": "Perfil objetivo, com atenção ao conteúdo literal do edital e aos conhecimentos específicos.",
        "foco": ["lei seca", "conteúdo específico", "português", "informática"],
        "alertas": ["exceções legais", "termos técnicos", "alternativas muito parecidas"],
        "fontes": ["https://www.institutoconsulplan.org.br/", "https://www.planalto.gov.br/"],
    },
    "Fundação Carlos Chagas": {
        "descricao": "Mesmo perfil da FCC; cadastro separado para facilitar a busca de editais e provas antigas.",
        "foco": ["legislação", "conceitos", "português", "conhecimentos específicos"],
        "alertas": ["prazos", "competências", "diferenças entre dispositivos"],
        "fontes": ["https://www.concursosfcc.com.br/", "https://www.planalto.gov.br/"],
    },
    "Legalle Concursos": {
        "descricao": "Banca regional com cobrança diretamente relacionada ao edital e à legislação indicada.",
        "foco": ["conteúdo do edital", "legislação municipal/estadual", "português", "conhecimentos específicos"],
        "alertas": ["normas locais", "atualização da lei", "detalhes do cargo"],
        "fontes": ["https://legalleconcursos.com.br/"],
    },
    "Fundatec": {
        "descricao": "Costuma cobrar o programa do edital com questões objetivas e legislação aplicada ao órgão.",
        "foco": ["legislação local", "conhecimentos específicos", "português", "administração pública"],
        "alertas": ["normas do órgão", "conceitos administrativos", "atualizações legislativas"],
        "fontes": ["https://www.fundatec.org.br/"],
    },
    "Quadrix": {
        "descricao": "Cobrança objetiva de legislação profissional, conhecimentos gerais e específicos conforme o conselho ou órgão.",
        "foco": ["legislação profissional", "ética", "conhecimentos específicos", "português"],
        "alertas": ["resoluções", "competências do conselho", "prazos e penalidades"],
        "fontes": ["https://www.quadrix.org.br/"],
    },
    "IADES": {"descricao": "Questões objetivas com peso relevante para o conteúdo específico e a legislação do edital.", "foco": ["conhecimentos específicos", "legislação", "português", "atualidades quando previstas"], "alertas": ["comandos negativos", "detalhes normativos", "conceitos próximos"], "fontes": ["https://www.iades.com.br/"]},
    "Consulplan": {"descricao": "Perfil objetivo e bastante ligado ao conteúdo programático do edital.", "foco": ["lei seca", "conhecimentos específicos", "português", "informática"], "alertas": ["exceções", "termos técnicos", "alternativas parecidas"], "fontes": ["https://www.institutoconsulplan.org.br/"]},
    "FUNDATEC": {"descricao": "Banca com forte presença regional e cobrança do edital e das normas do órgão.", "foco": ["legislação local", "específicos", "português", "administração pública"], "alertas": ["normas locais", "atualizações legais", "atribuições do cargo"], "fontes": ["https://www.fundatec.org.br/"]},
    "Instituto Verbena": {"descricao": "Cobra o programa do edital com questões objetivas e conteúdo específico do cargo.", "foco": ["específicos", "legislação", "português", "políticas públicas"], "alertas": ["normas locais", "conceitos técnicos", "comando da questão"], "fontes": ["https://institutoverbena.ufg.br/"]},
    "COPESE": {"descricao": "Perfil acadêmico e objetivo, com atenção ao conteúdo previsto no edital.", "foco": ["conhecimentos específicos", "português", "legislação", "raciocínio lógico"], "alertas": ["definições", "exceções", "interpretação"], "fontes": ["https://concursos.uft.edu.br/"]},
    "CESGRANRIO": {"descricao": "Questões contextualizadas, com interpretação, aplicação e cobrança equilibrada do conteúdo.", "foco": ["interpretação", "situações práticas", "português", "conhecimentos específicos"], "alertas": ["texto-base", "dados do enunciado", "alternativas próximas"], "fontes": ["https://www.cesgranrio.org.br/concursos/"]},
    "Instituto Mais": {"descricao": "Questões objetivas com cobrança do conteúdo programático e das normas indicadas.", "foco": ["legislação", "específicos", "português", "informática"], "alertas": ["detalhes do edital", "prazos", "atribuições"], "fontes": ["https://www.institutomais.org.br/"]},
    "IDECAN": {"descricao": "Cobra a base do edital combinando legislação, teoria e conhecimentos específicos.", "foco": ["específicos", "legislação", "português", "raciocínio lógico"], "alertas": ["exceções", "conceitos similares", "atualização normativa"], "fontes": ["https://www.idecan.org.br/"]},
    "Instituto ACCESS": {"descricao": "Perfil objetivo, com foco no conteúdo expressamente previsto no edital.", "foco": ["específicos", "legislação", "português", "informática"], "alertas": ["lei seca", "termos técnicos", "comandos negativos"], "fontes": ["https://concursos.access.org.br/"]},
    "IBGP": {"descricao": "Banca com cobrança objetiva e forte relação com o conteúdo programático, a legislação indicada e os conhecimentos específicos do cargo.", "foco": ["conhecimentos específicos", "legislação do edital", "língua portuguesa", "informática e raciocínio lógico quando previstos"], "alertas": ["detalhes da legislação", "competências e atribuições", "alternativas com conceitos próximos"], "fontes": ["https://www.ibgpconcursos.com.br/", "https://www.planalto.gov.br/"]},
    "Outra banca": {"descricao": "A banca será analisada a partir do edital enviado e das fontes indicadas pelo usuário.", "foco": ["conteúdo do edital", "legislação aplicável", "provas anteriores", "perfil das questões"], "alertas": ["confirmar a fonte oficial", "verificar a data do edital", "não confundir tendência com regra"], "fontes": ["https://www.gov.br/pt-br/servicos/consultar-concursos-publicos"]},
}

RAIO_X_DETALHES = {
    "IBGP": {
        "estilo": "Objetivo e direto, com enunciados geralmente ligados ao conteúdo do edital e aos conhecimentos específicos do cargo.",
        "como_cobra": ["lei seca e conceitos fundamentais", "interpretação de textos e comandos", "atribuições e competências do cargo", "legislação municipal, estadual ou do órgão quando prevista"],
        "formato": "Questões objetivas de múltipla escolha, com alternativas próximas e atenção ao comando exato.",
        "armadilhas": ["troca de conceitos semelhantes", "inversão de competência ou prazo", "alternativa correta apenas em parte", "resposta que extrapola o texto da norma"],
    },
    "CESGRANRIO": {
        "estilo": "Contextualizado e aplicado, com uso de situações práticas e textos-base.",
        "como_cobra": ["interpretação", "aplicação prática", "língua portuguesa", "conhecimentos específicos do cargo"],
        "formato": "Múltipla escolha, frequentemente exigindo leitura cuidadosa do enunciado e dos dados apresentados.",
        "armadilhas": ["ignorar informação do texto-base", "alternativas parcialmente corretas", "conclusão que não decorre do enunciado"],
    },
    "FGV": {
        "estilo": "Interpretativo e aplicado, com alternativas próximas e situações-problema.",
        "como_cobra": ["interpretação", "casos práticos", "conceitos e exceções", "aplicação da legislação"],
        "formato": "Múltipla escolha com enunciados contextualizados e comandos que podem exigir a alternativa incorreta.",
        "armadilhas": ["comando negativo", "regra confundida com exceção", "alternativa quase correta"],
    },
    "Cebraspe": {
        "estilo": "Analítico e atento a detalhes, normalmente no formato certo ou errado.",
        "como_cobra": ["literalidade da lei", "exceções", "prazos", "interpretação e aplicação do conteúdo"],
        "formato": "Itens de julgamento, em que uma palavra pode alterar toda a correção da afirmação.",
        "armadilhas": ["palavras absolutas", "troca de conceitos", "generalização indevida", "detalhe fora da regra"],
    },
}

def _detalhes_banca(banca: str, perfil: dict) -> dict:
    return RAIO_X_DETALHES.get(banca, {
        "estilo": perfil["descricao"],
        "como_cobra": perfil["foco"],
        "formato": "Questões objetivas conforme o edital e o padrão das provas anteriores disponíveis.",
        "armadilhas": perfil["alertas"],
    })

def _normalizar_raio_x(valor: str) -> str:
    return unicodedata.normalize("NFKD", (valor or "")).encode("ascii", "ignore").decode("ascii").casefold()

def _extrair_contexto_cargo(texto: str, cargo: str) -> tuple[str, bool, str]:
    """Limita a leitura ao trecho do cargo, em vez de pesquisar o edital inteiro."""
    linhas = [linha.strip() for linha in (texto or "").splitlines() if linha.strip()]
    if not linhas or not cargo.strip():
        return texto or "", False, ""
    cargo_normalizado = re.sub(r"\s+", " ", _normalizar_raio_x(cargo)).strip()
    tokens = [token for token in re.findall(r"[a-z0-9]+", cargo_normalizado) if len(token) > 2 and token not in {"de", "da", "do", "das", "dos"}]
    if not tokens:
        return texto or "", False, ""
    melhor = None
    maior = 0
    # Primeiro tenta a expressão completa, evitando confundir cargos que
    # possuem somente uma palavra em comum, como Farmácia e Endemias.
    for indice, linha in enumerate(linhas):
        normalizada = re.sub(r"\s+", " ", _normalizar_raio_x(linha)).strip()
        if cargo_normalizado in normalizada:
            melhor = indice
            maior = len(tokens)
            break
    # Fallback para abreviações presentes no edital (ex.: ACE).
    for indice, linha in enumerate(linhas):
        if melhor is not None:
            break
        normalizada = _normalizar_raio_x(linha)
        pontuacao = sum(1 for token in tokens if token in normalizada)
        if pontuacao > maior:
            maior = pontuacao
            melhor = indice
    if melhor is None or maior < max(1, int(len(tokens) * 0.6)):
        return texto or "", False, ""
    inicio = max(0, melhor - 5)
    fim = min(len(linhas), melhor + 90)
    # Não atravessa o começo de outro cargo/área do conteúdo programático.
    for indice in range(melhor + 8, fim):
        linha_original = linhas[indice]
        linha = _normalizar_raio_x(linhas[indice])
        if re.match(r"^(cargo|funcao|funcao publica|conteudo programatico)\s*[:\-]", linha) or (linha_original.isupper() and len(linha_original.split()) <= 12):
            fim = indice
            break
    return "\n".join(linhas[inicio:fim]), True, linhas[melhor]

def _extrair_contexto_geral(texto: str) -> str:
    """Lê somente blocos explicitamente comuns a todos os cargos."""
    linhas = [linha.strip() for linha in (texto or "").splitlines() if linha.strip()]
    marcadores = ("conhecimentos gerais", "conhecimentos básicos", "conteúdo comum", "prova de conhecimentos gerais")
    trechos = []
    for indice, linha in enumerate(linhas):
        normalizada = _normalizar_raio_x(linha)
        if any(_normalizar_raio_x(marcador) in normalizada for marcador in marcadores):
            trechos.extend(linhas[max(0, indice - 2):min(len(linhas), indice + 90)])
    return "\n".join(dict.fromkeys(trechos))

def _extrair_materias_edital(texto: str) -> dict:
    """Extrai matérias somente do trecho selecionado do cargo."""
    padroes = [
        ("Língua Portuguesa", "geral"), ("Raciocínio Lógico", "geral"), ("Matemática", "geral"),
        ("Informática", "geral"), ("Atualidades", "geral"), ("Conhecimentos Gerais", "geral"),
        ("Direito Constitucional", "especifica"), ("Direito Administrativo", "especifica"),
        ("Direito Penal", "especifica"), ("Direito Processual Penal", "especifica"), ("Direito Civil", "especifica"),
        ("Administração Pública", "especifica"), ("Legislação", "especifica"),
        ("Conhecimentos Específicos", "especifica"), ("Ética", "especifica"),
        ("Redação Oficial", "geral"), ("Arquivologia", "especifica"), ("Contabilidade", "especifica"),
        ("Economia", "especifica"), ("Administração Financeira e Orçamentária", "especifica"),
        ("Direitos Humanos", "especifica"), ("Medicina Legal", "especifica"),
        ("Saúde Pública", "especifica"), ("Sistema Único de Saúde (SUS)", "especifica"),
        ("Vigilância Epidemiológica", "especifica"), ("Endemias", "especifica"), ("Dengue", "especifica"),
        ("Zoonoses", "especifica"), ("Educação em Saúde", "especifica"), ("Leishmaniose", "especifica"),
    ]
    normalizado = _normalizar_raio_x(texto)
    gerais, especificas = [], []
    for materia, grupo in padroes:
        if _normalizar_raio_x(materia) in normalizado:
            (gerais if grupo == "geral" else especificas).append(materia)
    return {"gerais": gerais, "especificas": especificas, "todas": gerais + especificas}

@app.get("/estudos/raio-x", response_class=HTMLResponse)
async def view_raio_x_banca(request: Request, analise_id: int = None):
    user_id = request.cookies.get("user_id")
    if not user_id:
        return RedirectResponse(url="/estudos/login")
    usuario_id = int(user_id)
    with Session(engine) as session:
        historico = session.exec(select(EstudoRaioXBanca).where(EstudoRaioXBanca.usuario_id == usuario_id).order_by(desc(EstudoRaioXBanca.criado_em))).all()
        analise = session.get(EstudoRaioXBanca, analise_id) if analise_id else (historico[0] if historico else None)
        if analise and analise.usuario_id != usuario_id:
            analise = None
    dados = None
    if analise:
        materias_salvas = json.loads(analise.materias_json or "{}")
        if isinstance(materias_salvas, list):
            materias_salvas = {"gerais": materias_salvas, "especificas": [], "todas": materias_salvas}
        dados = {
            "materias": materias_salvas,
            "analise": json.loads(analise.analise_json or "{}"),
            "fontes": json.loads(analise.fontes_json or "[]"),
        }
        perfil_atual = RAIO_X_BANCAS.get(analise.banca) or RAIO_X_BANCAS["Outra banca"]
        detalhes_atual = _detalhes_banca(analise.banca, perfil_atual)
        dados["analise"].setdefault("estilo", detalhes_atual["estilo"])
        dados["analise"].setdefault("como_cobra", detalhes_atual["como_cobra"])
        dados["analise"].setdefault("formato", detalhes_atual["formato"])
        dados["analise"].setdefault("armadilhas_banca", detalhes_atual["armadilhas"])
    return templates.TemplateResponse(request=request, name="estudos_raio_x.html", context={"config": APP_CONFIG, "bancas": sorted(RAIO_X_BANCAS), "historico": historico, "analise": analise, "dados": dados})

@app.post("/estudos/raio-x")
async def criar_raio_x_banca(request: Request, banca: str = Form(...), cargo: str = Form(""), edital: UploadFile = File(...)):
    user_id = request.cookies.get("user_id")
    if not user_id:
        return RedirectResponse(url="/estudos/login", status_code=303)
    perfil = RAIO_X_BANCAS.get(banca) or RAIO_X_BANCAS["Outra banca"]
    if not edital.filename or not edital.filename.lower().endswith(".pdf"):
        return RedirectResponse(url="/estudos/raio-x?erro=pdf", status_code=303)
    conteudo = await edital.read()
    if len(conteudo) > 15 * 1024 * 1024:
        return RedirectResponse(url="/estudos/raio-x?erro=tamanho", status_code=303)
    texto = ""
    try:
        from io import BytesIO
        from pypdf import PdfReader
        reader = PdfReader(BytesIO(conteudo))
        texto = "\n".join((pagina.extract_text() or "") for pagina in reader.pages)
    except Exception:
        texto = ""
    cargo_texto, cargo_encontrado, linha_cargo = _extrair_contexto_cargo(texto, cargo)
    materias = _extrair_materias_edital(cargo_texto)
    # Muitos editais colocam Português e Raciocínio Lógico em uma seção comum,
    # fora do bloco do cargo. Complementa somente matérias gerais, sem puxar
    # conhecimentos específicos de outros cargos.
    materias_gerais_comuns = _extrair_materias_edital(_extrair_contexto_geral(texto))
    gerais_permitidas = {"Língua Portuguesa", "Raciocínio Lógico", "Matemática", "Informática", "Atualidades", "Conhecimentos Gerais", "Redação Oficial"}
    for materia in materias_gerais_comuns.get("gerais", []):
        if materia in gerais_permitidas and materia not in materias["gerais"]:
            materias["gerais"].append(materia)
    # Português e Raciocínio Lógico também podem aparecer em um bloco
    # separado do título "Conhecimentos Gerais"; são complementados somente
    # quando constam no texto do edital, sem fazer isso para Informática.
    materias_gerais_texto = _extrair_materias_edital(texto).get("gerais", [])
    for materia in ("Língua Portuguesa", "Raciocínio Lógico"):
        if materia in materias_gerais_texto and materia not in materias["gerais"]:
            materias["gerais"].append(materia)
    materias["todas"] = materias["gerais"] + materias["especificas"]
    materias_normalizadas = {_normalizar_raio_x(materia) for materia in materias["todas"]}
    foco_filtrado = []
    for item_foco in perfil["foco"]:
        texto_foco = item_foco
        if "informatica" in _normalizar_raio_x(texto_foco) and "informatica" not in materias_normalizadas:
            texto_foco = re.sub(r"informática\s*(?:e\s*)?", "", texto_foco, flags=re.IGNORECASE).strip()
            texto_foco = re.sub(r"\s+", " ", texto_foco)
        if "raciocinio logico" in _normalizar_raio_x(texto_foco) and "raciocinio logico" not in materias_normalizadas:
            texto_foco = re.sub(r"raciocínio lógico(?:\s+quando previstos)?", "", texto_foco, flags=re.IGNORECASE).strip(" ,;-")
        if texto_foco and texto_foco not in foco_filtrado:
            foco_filtrado.append(texto_foco)
    termo_busca = f"{banca} {cargo.strip() or 'concurso'} prova anterior"
    busca_provas = f"https://www.google.com/search?q={quote(termo_busca, safe='')}"
    detalhes = _detalhes_banca(banca, perfil)
    analise = {
        "descricao": perfil["descricao"],
        "foco": foco_filtrado,
        "estilo": detalhes["estilo"],
        "como_cobra": detalhes["como_cobra"],
        "formato": detalhes["formato"],
        "armadilhas_banca": detalhes["armadilhas"],
        "alertas": perfil["alertas"],
        "limites": ["As matérias foram procuradas no trecho do cargo informado, não no edital inteiro.", "O edital confirma o conteúdo; ele não prova sozinho a frequência de cada tema.", "Percentuais e ranking só serão calculados quando houver questões/provas cadastradas.", "As tendências devem ser revisadas quando sair novo edital."],
        "cargo_encontrado": cargo_encontrado,
        "linha_cargo": linha_cargo,
        "escopo": "Trecho específico do cargo" if cargo_encontrado else "Edital inteiro: o cargo não foi localizado no texto extraído",
        "referencia_provas": busca_provas,
        "texto_extraido": bool(texto.strip()),
        "paginas": len(getattr(reader, "pages", [])) if texto else 0,
    }
    with Session(engine) as session:
        # Mantém somente a análise atual por usuário, evitando crescimento indefinido.
        antigas = session.exec(select(EstudoRaioXBanca).where(EstudoRaioXBanca.usuario_id == int(user_id))).all()
        for antiga in antigas:
            session.delete(antiga)
        item = EstudoRaioXBanca(usuario_id=int(user_id), banca=banca, cargo=cargo.strip() or None, edital_nome=edital.filename, edital_texto=cargo_texto[:200000], materias_json=json.dumps(materias, ensure_ascii=False), analise_json=json.dumps(analise, ensure_ascii=False), fontes_json=json.dumps(perfil["fontes"] + [busca_provas], ensure_ascii=False))
        session.add(item)
        session.commit()
        session.refresh(item)
        item_id = item.id
    return RedirectResponse(url=f"/estudos/raio-x?analise_id={item_id}", status_code=303)

@app.post("/estudos/raio-x/excluir")
async def excluir_raio_x_banca(request: Request):
    user_id = request.cookies.get("user_id")
    if not user_id:
        return RedirectResponse(url="/estudos/login", status_code=303)
    with Session(engine) as session:
        analises = session.exec(select(EstudoRaioXBanca).where(EstudoRaioXBanca.usuario_id == int(user_id))).all()
        for analise in analises:
            session.delete(analise)
        session.commit()
    return RedirectResponse(url="/estudos/raio-x", status_code=303)

@app.get("/estudos/grafico-semanal")
async def estudos_grafico_semanal(request: Request, semana: str = None, assunto: str = None):
    user_id = request.cookies.get("user_id")
    if not user_id:
        raise HTTPException(status_code=401)
    try:
        usuario_id = int(user_id)
    except ValueError:
        raise HTTPException(status_code=401)

    hoje = datetime.now().date()
    inicio_atual = hoje - timedelta(days=hoje.weekday())
    try:
        data_semana = datetime.strptime(semana, "%Y-%m-%d").date() if semana else inicio_atual
    except ValueError:
        data_semana = inicio_atual
    semana_inicio = data_semana - timedelta(days=data_semana.weekday())
    semana_fim = semana_inicio + timedelta(days=6)
    filtro = (assunto or "").strip().lower()

    with Session(engine) as session:
        registros = session.exec(
            select(EstudoSessao).where(
                EstudoSessao.usuario_id == usuario_id,
                EstudoSessao.data_estudo >= semana_inicio,
                EstudoSessao.data_estudo <= semana_fim,
            )
        ).all()

    if filtro:
        registros = [r for r in registros if filtro in (r.assunto or "").lower()]
    nomes_dias = ["Seg", "Ter", "Qua", "Qui", "Sex", "Sáb", "Dom"]
    dias = []
    for indice, nome in enumerate(nomes_dias):
        data = semana_inicio + timedelta(days=indice)
        dias.append({
            "dia": nome,
            "data": data.isoformat(),
            "questoes": sum((r.quantidade_questoes or 0) for r in registros if r.data_estudo == data),
        })
    return {
        "inicio": semana_inicio.isoformat(),
        "fim": semana_fim.isoformat(),
        "semana_atual": inicio_atual.isoformat(),
        "anterior": (semana_inicio - timedelta(days=7)).isoformat(),
        "proxima": (semana_inicio + timedelta(days=7)).isoformat(),
        "dias": dias,
    }


@app.get("/estudos/progresso")
async def estudos_progresso(request: Request, data_progresso: str = None):
    user_id = request.cookies.get("user_id")
    if not user_id:
        raise HTTPException(status_code=401)
    try:
        usuario_id = int(user_id)
    except ValueError:
        raise HTTPException(status_code=401)
    hoje = datetime.now().date()
    try:
        data = datetime.strptime(data_progresso, "%Y-%m-%d").date() if data_progresso else hoje
    except ValueError:
        data = hoje
    with Session(engine) as session:
        registros = session.exec(select(EstudoSessao).where(
            EstudoSessao.usuario_id == usuario_id,
            EstudoSessao.data_estudo == data,
        )).all()
    por_materia = {}
    for registro in registros:
        materia = (registro.assunto or "Geral").strip()
        item = por_materia.setdefault(materia, {"questoes": 0, "acertos": 0})
        item["questoes"] += registro.quantidade_questoes or 0
        item["acertos"] += registro.total_acertos or 0
    itens = [{
        "materia": materia,
        "questoes": valores["questoes"],
        "acertos": valores["acertos"],
        "percentual": min(100, round((valores["acertos"] / valores["questoes"]) * 100, 1)) if valores["questoes"] else 0,
    } for materia, valores in sorted(por_materia.items(), key=lambda item: item[1]["questoes"], reverse=True)]
    return {"data": data.isoformat(), "label": "Hoje" if data == hoje else ("Ontem" if data == hoje - timedelta(days=1) else data.strftime("%d/%m/%Y")), "itens": itens}


@app.get("/estudos", response_class=HTMLResponse)
async def view_estudos(request: Request, success: str = None, assunto: str = None, data_inicio: str = None, data_fim: str = None, semana: str = None, editar_id: int = None, data_progresso: str = None):
    user_id = request.cookies.get("user_id")
    if not user_id:
        return RedirectResponse(url="/login")

    try:
        usuario_id = int(user_id)
    except ValueError:
        return RedirectResponse(url="/login")

    hoje = datetime.now().date()
    ontem = hoje - timedelta(days=1)
    inicio_semana = hoje - timedelta(days=hoje.weekday())
    try:
        data_semana = datetime.strptime(semana, "%Y-%m-%d").date() if semana else inicio_semana
        semana_inicio = data_semana - timedelta(days=data_semana.weekday())
    except ValueError:
        semana_inicio = inicio_semana
    fim_semana = semana_inicio + timedelta(days=6)
    inicio_dia = datetime.combine(hoje, datetime.min.time())
    inicio_proximo_dia = inicio_dia + timedelta(days=1)

    with Session(engine) as session:
        registros = session.exec(
            select(EstudoSessao).where(EstudoSessao.usuario_id == usuario_id).order_by(desc(EstudoSessao.data_estudo), desc(EstudoSessao.criado_em))
        ).all()

        planos = session.exec(
            select(EstudoPlano).where(
                EstudoPlano.usuario_id == usuario_id,
                EstudoPlano.semana_inicio == semana_inicio,
            ).order_by(EstudoPlano.materia)
        ).all()
        has_grid = any(p.slot is not None for p in planos)
        if not has_grid:
            todas = session.exec(
                select(EstudoPlano.semana_inicio)
                .where(
                    EstudoPlano.usuario_id == usuario_id,
                    EstudoPlano.semana_inicio < semana_inicio,
                    EstudoPlano.slot.isnot(None)
                )
                .order_by(desc(EstudoPlano.semana_inicio))
            ).all()
            if todas:
                ultima_semana = todas[0]
                planos_antigos = session.exec(
                    select(EstudoPlano).where(
                        EstudoPlano.usuario_id == usuario_id,
                        EstudoPlano.semana_inicio == ultima_semana,
                        EstudoPlano.slot.isnot(None)
                    ).order_by(EstudoPlano.materia)
                ).all()
                planos.extend(planos_antigos)
        
        if not planos:
            todas_simples = session.exec(
                select(EstudoPlano.semana_inicio)
                .where(EstudoPlano.usuario_id == usuario_id, EstudoPlano.semana_inicio < semana_inicio)
                .order_by(desc(EstudoPlano.semana_inicio))
            ).all()
            if todas_simples:
                planos = session.exec(
                    select(EstudoPlano).where(
                        EstudoPlano.usuario_id == usuario_id,
                        EstudoPlano.semana_inicio == todas_simples[0]
                    ).order_by(EstudoPlano.materia)
                ).all()
        editar = session.get(EstudoSessao, editar_id) if editar_id else None
        if editar and editar.usuario_id != usuario_id:
            editar = None

    filtro = (assunto or "").strip().lower()
    try:
        data_progresso_selecionada = datetime.strptime(data_progresso, "%Y-%m-%d").date() if data_progresso else hoje
    except ValueError:
        data_progresso_selecionada = hoje
    try:
        filtro_data_inicio = datetime.strptime(data_inicio, "%Y-%m-%d").date() if data_inicio else None
    except ValueError:
        filtro_data_inicio = None
    try:
        filtro_data_fim = datetime.strptime(data_fim, "%Y-%m-%d").date() if data_fim else None
    except ValueError:
        filtro_data_fim = None
    registros_historico = [
        registro for registro in registros
        if (not filtro or filtro in (registro.assunto or "").lower())
        and (not filtro_data_inicio or registro.data_estudo >= filtro_data_inicio)
        and (not filtro_data_fim or registro.data_estudo <= filtro_data_fim)
    ]

    # Os cards principais representam o total geral de todos os estudos.
    registros_calculo = registros
    total_questoes = sum((r.quantidade_questoes or 0) for r in registros_calculo)
    total_acertos = sum((r.total_acertos or 0) for r in registros_calculo)
    total_erros = sum((r.total_erros or 0) for r in registros_calculo)
    registros_simulados = [r for r in registros_calculo if getattr(r, "tipo_registro", "estudo") == "simulado"]
    registros_estudos = [r for r in registros_calculo if getattr(r, "tipo_registro", "estudo") != "simulado"]
    questoes_simulados = sum((r.quantidade_questoes or 0) for r in registros_simulados)
    acertos_simulados = sum((r.total_acertos or 0) for r in registros_simulados)
    questoes_estudos = sum((r.quantidade_questoes or 0) for r in registros_estudos)
    acertos_estudos = sum((r.total_acertos or 0) for r in registros_estudos)
    aproveitamento_simulados = min(100, round((acertos_simulados / questoes_simulados) * 100, 1)) if questoes_simulados else 0
    aproveitamento_estudos = min(100, round((acertos_estudos / questoes_estudos) * 100, 1)) if questoes_estudos else 0
    def calc_segundos(r):
        if r.horario_inicio and r.horario_fim:
            try:
                hi = datetime.strptime(r.horario_inicio, "%H:%M")
                hf = datetime.strptime(r.horario_fim, "%H:%M")
                diff = (hf - hi).total_seconds()
                if diff < 0: diff += 24 * 3600 # cruzou meia noite
                return int(diff)
            except ValueError:
                pass
        return 0

    total_cronometro_segundos = sum(calc_segundos(r) for r in registros)
    cronometro_hoje_segundos = sum(calc_segundos(r) for r in registros if r.data_estudo == hoje)
    horas_cronometro_total = f"{total_cronometro_segundos // 3600}h {(total_cronometro_segundos % 3600) // 60:02d}min"
    horas_cronometro_hoje = f"{cronometro_hoje_segundos // 3600}h {(cronometro_hoje_segundos % 3600) // 60:02d}min"

    semana_questoes = sum(
        (r.quantidade_questoes or 0)
        for r in registros
        if r.data_estudo and semana_inicio <= r.data_estudo <= fim_semana
    )
    meta_diaria = 40
    meta_semanal = 200
    total_hoje = sum(
        (r.quantidade_questoes or 0)
        for r in registros
        if r.data_estudo == hoje
    )
    progresso_diario = min(100, int((total_hoje / meta_diaria) * 100)) if meta_diaria else 0
    progresso_semanal = min(100, int((semana_questoes / meta_semanal) * 100)) if meta_semanal else 0

    assuntos = {}
    for r in registros:
        assunto_key = r.assunto or "Geral"
        assuntos[assunto_key] = assuntos.get(assunto_key, 0) + (r.quantidade_questoes or 0)
    assuntos_ordenados = sorted(assuntos.items(), key=lambda item: item[1], reverse=True)[:5]

    grafico_dias = []
    nomes_dias = ["Seg", "Ter", "Qua", "Qui", "Sex", "Sáb", "Dom"]
    for indice, nome_dia in enumerate(nomes_dias):
        data_grafico = semana_inicio + timedelta(days=indice)
        quantidade = sum(
            (registro.quantidade_questoes or 0)
            for registro in registros
            if registro.data_estudo == data_grafico
        )
        grafico_dias.append({"dia": nome_dia, "data": data_grafico, "questoes": quantidade})

    desempenho_materias = {}
    for registro in registros:
        materia = (registro.assunto or "Geral").strip()
        item = desempenho_materias.setdefault(materia, {"questoes": 0, "acertos": 0})
        item["questoes"] += registro.quantidade_questoes or 0
        item["acertos"] += registro.total_acertos or 0
    grafico_materias = [
        {
            "materia": materia,
            "questoes": valores["questoes"],
            "acertos": valores["acertos"],
            "percentual": min(100, round((valores["acertos"] / valores["questoes"]) * 100, 1)) if valores["questoes"] else 0,
        }
        for materia, valores in sorted(desempenho_materias.items(), key=lambda item: item[1]["questoes"], reverse=True)
    ][:8]
    
    desempenho_materias_hoje = {}
    for registro in registros:
        if registro.data_estudo == data_progresso_selecionada:
            materia = (registro.assunto or "Geral").strip()
            item = desempenho_materias_hoje.setdefault(materia, {"questoes": 0, "acertos": 0})
            item["questoes"] += registro.quantidade_questoes or 0
            item["acertos"] += registro.total_acertos or 0
    grafico_materias_hoje = [
        {
            "materia": materia,
            "questoes": valores["questoes"],
            "acertos": valores["acertos"],
            "percentual": min(100, round((valores["acertos"] / valores["questoes"]) * 100, 1)) if valores["questoes"] else 0,
        }
        for materia, valores in sorted(desempenho_materias_hoje.items(), key=lambda item: item[1]["questoes"], reverse=True)
    ]
    if data_progresso_selecionada == hoje:
        label_progresso = "Hoje"
    elif data_progresso_selecionada == ontem:
        label_progresso = "Ontem"
    else:
        label_progresso = data_progresso_selecionada.strftime("%d/%m/%Y")
    max_questoes_grafico = max((item["questoes"] for item in grafico_dias), default=1) or 1
    materias_com_dados = [item for item in grafico_materias if item["questoes"] > 0]
    melhor_materia = max(materias_com_dados, key=lambda item: item["percentual"]) if materias_com_dados else None
    pior_materia = min(materias_com_dados, key=lambda item: item["percentual"]) if materias_com_dados else None
    materia_melhorar = pior_materia

    realizados_por_materia = {}
    for registro in registros:
        if registro.data_estudo and semana_inicio <= registro.data_estudo <= fim_semana:
            chave = (registro.assunto or "Estudo").strip().lower()
            realizados_por_materia[chave] = realizados_por_materia.get(chave, 0) + (registro.quantidade_questoes or 0)
    plano_comparativo = [
        {
            "materia": plano.materia,
            "dias_semana": plano.dias_semana,
            "meta_questoes": plano.meta_questoes,
            "realizadas": realizados_por_materia.get(plano.materia.strip().lower(), 0),
            "percentual": min(100, int((realizados_por_materia.get(plano.materia.strip().lower(), 0) / plano.meta_questoes) * 100)) if plano.meta_questoes else 0,
            "dia_semana": plano.dia_semana,
            "horario_inicio": plano.horario_inicio,
            "horario_fim": plano.horario_fim,
        }
        for plano in planos
    ]
    total_planejado = sum(item["meta_questoes"] for item in plano_comparativo)
    total_planejado_realizado = sum(item["realizadas"] for item in plano_comparativo)
    progresso_plano = min(100, int((total_planejado_realizado / total_planejado) * 100)) if total_planejado else 0
    dias_grade = ["Segunda", "Terça", "Quarta", "Quinta", "Sexta", "Sábado", "Domingo"]
    horarios_grade = [("21:00", "22:00"), ("22:15", "00:15"), ("00:31", "01:15")]
    for plano in planos:
        if plano.slot is not None and 0 <= plano.slot < len(horarios_grade) and plano.horario_inicio and plano.horario_fim:
            horarios_grade[plano.slot] = (plano.horario_inicio, plano.horario_fim)
    grade = {
        f"{item['horario_inicio']}-{item['horario_fim']}-{item['dia_semana']}": item["materia"]
        for item in plano_comparativo
        if item["dia_semana"] and item["horario_inicio"] and item["horario_fim"]
    }
    materias_unicas = {}
    for materia in grade.values():
        if not materia:
            continue
        chave = unicodedata.normalize("NFKD", materia).encode("ascii", "ignore").decode("ascii").casefold()
        chave = re.sub(r"[^a-z0-9]+", " ", chave).strip()
        chave = re.sub(r"\bconh\.?\b|\bconhec\.?\b", "conhecimentos", chave)
        materias_unicas.setdefault(chave, materia.strip())
    for plano in planos:
        if not plano.materia:
            continue
        chave = unicodedata.normalize("NFKD", plano.materia).encode("ascii", "ignore").decode("ascii").casefold()
        chave = re.sub(r"[^a-z0-9]+", " ", chave).strip()
        chave = re.sub(r"\bconh\.?\b|\bconhec\.?\b", "conhecimentos", chave)
        materias_unicas.setdefault(chave, plano.materia.strip())
    for registro in registros:
        if not registro.assunto:
            continue
        chave = unicodedata.normalize("NFKD", registro.assunto).encode("ascii", "ignore").decode("ascii").casefold()
        chave = re.sub(r"[^a-z0-9]+", " ", chave).strip()
        chave = re.sub(r"\bconh\.?\b|\bconhec\.?\b", "conhecimentos", chave)
        materias_unicas.setdefault(chave, registro.assunto.strip())
    materias_grade = sorted(materias_unicas.values(), key=str.casefold)
    paleta_materias = [
        ("#38bdf8", "rgba(56, 189, 248, 0.12)"),
        ("#fbbf24", "rgba(251, 191, 36, 0.12)"),
        ("#a3e635", "rgba(163, 230, 53, 0.12)"),
        ("#fb7185", "rgba(251, 113, 133, 0.12)"),
        ("#c084fc", "rgba(192, 132, 252, 0.12)"),
        ("#2dd4bf", "rgba(45, 212, 191, 0.12)"),
        ("#fb923c", "rgba(251, 146, 60, 0.12)"),
        ("#818cf8", "rgba(129, 140, 248, 0.12)"),
    ]
    def normalizar_chave_materia(materia):
        chave = unicodedata.normalize("NFKD", materia).encode("ascii", "ignore").decode("ascii").casefold()
        chave = re.sub(r"[^a-z0-9]+", " ", chave).strip()
        return re.sub(r"\bconh\.?\b|\bconhec\.?\b", "conhecimentos", chave)

    chaves_materias_quadro = sorted({normalizar_chave_materia(materia) for materia in grade.values() if materia and materia.strip()})
    materias_quadro = [materias_unicas[chave] for chave in chaves_materias_quadro]
    cores_por_chave = {
        chave: paleta_materias[indice % len(paleta_materias)]
        for indice, chave in enumerate(chaves_materias_quadro)
    }
    materia_cores = {
        materia.casefold(): cores_por_chave[normalizar_chave_materia(materia)]
        for materia in grade.values()
        if materia and materia.strip()
    }

    return templates.TemplateResponse(
        request=request,
        name="estudos.html",
        context={
            "config": APP_CONFIG,
            "registros": registros_historico,
            "total_questoes": total_questoes,
            "total_acertos": total_acertos,
            "total_erros": total_erros,
            "questoes_estudos": questoes_estudos,
            "acertos_estudos": acertos_estudos,
            "aproveitamento_estudos": aproveitamento_estudos,
            "questoes_simulados": questoes_simulados,
            "acertos_simulados": acertos_simulados,
            "aproveitamento_simulados": aproveitamento_simulados,
            "horas_cronometro_total": horas_cronometro_total,
            "horas_cronometro_hoje": horas_cronometro_hoje,
            "success": success,
            "meta_diaria": meta_diaria,
            "meta_semanal": meta_semanal,
            "total_hoje": total_hoje,
            "semana_questoes": semana_questoes,
            "progresso_diario": progresso_diario,
            "progresso_semanal": progresso_semanal,
            "assuntos_ordenados": assuntos_ordenados,
            "assunto_filtro": filtro,
            "data_inicio_filtro": data_inicio or "",
            "data_fim_filtro": data_fim or "",
            "planos": plano_comparativo,
            "semana_inicio": semana_inicio,
            "semana_fim": fim_semana,
            "semana_atual_inicio": inicio_semana,
            "semana_anterior": semana_inicio - timedelta(days=7),
            "semana_proxima": semana_inicio + timedelta(days=7),
            "data_hoje": hoje,
            "data_ontem": ontem,
            "total_planejado": total_planejado,
            "total_planejado_realizado": total_planejado_realizado,
            "progresso_plano": progresso_plano,
            "dias_grade": dias_grade,
            "horarios_grade": horarios_grade,
            "grade": grade,
            "materias_grade": materias_grade,
            "materias_quadro": materias_quadro,
            "materia_cores": materia_cores,
            "editar": editar,
            "grafico_dias": grafico_dias,
            "grafico_materias": grafico_materias,
            "grafico_materias_hoje": grafico_materias_hoje,
            "data_progresso_filtro": data_progresso_selecionada.isoformat() if data_progresso_selecionada else hoje.isoformat(),
            "label_progresso": label_progresso,
            "max_questoes_grafico": max_questoes_grafico,
            "melhor_materia": melhor_materia,
            "pior_materia": pior_materia,
            "materia_melhorar": materia_melhorar,
        }
    )

@app.post("/estudos/cronometro/concluir")
async def concluir_cronometro(request: Request):
    user_id = request.cookies.get("user_id")
    if not user_id:
        raise HTTPException(status_code=401, detail="Autenticação necessária")
    try:
        usuario_id = int(user_id)
        payload = await request.json()
    except (ValueError, TypeError):
        raise HTTPException(status_code=400, detail="Dados inválidos")

    if not isinstance(payload, dict):
        raise HTTPException(status_code=400, detail="Dados inválidos")
    session_id = payload.get("session_id")
    duracao_segundos = payload.get("duracao_segundos")
    materia = payload.get("materia")
    if not isinstance(session_id, str) or not session_id.strip() or len(session_id) > 64:
        raise HTTPException(status_code=400, detail="Identificador de sessão inválido")
    if isinstance(duracao_segundos, bool) or not isinstance(duracao_segundos, int) or not 60 <= duracao_segundos <= 36000:
        raise HTTPException(status_code=400, detail="Duração de sessão inválida")
    if materia is not None and not isinstance(materia, str):
        raise HTTPException(status_code=400, detail="Matéria inválida")

    with Session(engine) as session:
        existente = session.exec(
            select(EstudoCronometro).where(EstudoCronometro.session_id == session_id)
        ).first()
        if existente:
            return {"ok": True, "registrada": False}
        session.add(EstudoCronometro(
            usuario_id=usuario_id,
            session_id=session_id,
            materia=materia.strip()[:200] if materia else None,
            duracao_segundos=duracao_segundos,
        ))
        session.commit()
    return {"ok": True, "registrada": True}

@app.get("/estudos/modelo-excel")
async def download_modelo_excel(request: Request):
    user_id = request.cookies.get("user_id")
    if not user_id:
        return RedirectResponse(url="/login")

    import io
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Quadro de Horarios"

    colunas = ["Hora", "Segunda", "Terça", "Quarta", "Quinta", "Sexta", "Sábado", "Domingo"]
    ws.append(colunas)

    linhas = [
        ["21:00 às 22:00", "Português", "Raciocínio Lógico", "Informática", "Direito Constitucional", "Direito Administrativo", "Revisão Geral", "Simulado"],
        ["22:15 às 00:15", "Conhecimentos Específicos", "Conhecimentos Específicos", "Informática", "Português", "Conhecimentos Específicos", "Questões", "Descanso"],
        ["00:31 às 01:15", "Questões", "Questões", "Questões", "Questões", "Questões", "Flashcards", "Planejamento"],
    ]
    for r in linhas:
        ws.append(r)

    header_fill = PatternFill(start_color="0284C7", end_color="0284C7", fill_type="solid")
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Arial", size=10)
    center_align = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1")
    )

    for col_idx in range(1, len(colunas) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        cell.border = thin_border

    for row_idx in range(2, len(linhas) + 2):
        for col_idx in range(1, len(colunas) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = data_font
            cell.alignment = center_align
            cell.border = thin_border
            if row_idx % 2 == 1:
                cell.fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    larguras = {"A": 18, "B": 24, "C": 24, "D": 24, "E": 24, "F": 24, "G": 20, "H": 20}
    for col_letter, width in larguras.items():
        ws.column_dimensions[col_letter].width = width

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return Response(
        content=buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="modelo_quadro_estudos.xlsx"'}
    )


@app.post("/estudos/importar-plano", response_class=HTMLResponse)
async def importar_plano_estudo(request: Request, semana_inicio: str = Form(...)):
    user_id = request.cookies.get("user_id")
    if not user_id:
        return RedirectResponse(url="/login")

    try:
        usuario_id = int(user_id)
        data_inicio = datetime.strptime(semana_inicio, "%Y-%m-%d").date()
    except (ValueError, TypeError):
        return RedirectResponse(url="/estudos?success=erro", status_code=303)

    form = await request.form()
    texto_plano = str(form.get("plano_texto") or "").strip()
    arquivo = form.get("plano_arquivo")

    conteudo_bytes = None
    nome_arquivo = ""
    if arquivo is not None and getattr(arquivo, "filename", None):
        try:
            conteudo_bytes = await arquivo.read()
            nome_arquivo = (arquivo.filename or "").lower()
        except Exception:
            conteudo_bytes = None

    dias_semana_oficiais = [
        ("segunda", "Segunda"),
        ("terca", "Terça"),
        ("quarta", "Quarta"),
        ("quinta", "Quinta"),
        ("sexta", "Sexta"),
        ("sabado", "Sábado"),
        ("domingo", "Domingo"),
    ]
    dias_ordenados = ["Segunda", "Terça", "Quarta", "Quinta", "Sexta", "Sábado", "Domingo"]
    default_slots = [("21:00", "22:00"), ("22:15", "00:15"), ("00:31", "01:15")]

    palavras_ignoradas = {
        "id", "cod", "codigo", "código", "data", "hora", "horario", "horário", "horas",
        "tempo", "status", "questoes", "questões", "acertos", "erros", "obs", "observacao",
        "observação", "total", "dia", "dias", "semana", "materia", "materias", "matéria",
        "matérias", "disciplina", "disciplinas", "assunto", "assuntos", "conteudo",
        "conteúdo", "conteudos", "conteúdos", "ciclo", "fila", "nome", "descricao",
        "descrição", "ordem", "meta", "tipo", "item", "null", "none", "nan", "ok", "sim",
        "nao", "não",
    }

    def normalizar_materia(valor) -> str:
        if valor is None:
            return ""
        texto = re.sub(r"\s+", " ", str(valor)).strip(" -:;|/\n\t\r")
        return texto.strip()

    def eh_cabecalho_ou_invalido(texto: str) -> bool:
        if not texto:
            return True
        t_clean = unicodedata.normalize("NFKD", texto).encode("ascii", "ignore").decode("ascii").casefold().strip(" :;-#*")
        if t_clean in palavras_ignoradas:
            return True
        if re.fullmatch(r"^\d+([.,]\d+)?$", t_clean):
            return True
        if re.fullmatch(r"^\d{1,2}:\d{2}(\s*(?:-|–|a|ate|até)\s*\d{1,2}:\d{2})?$", t_clean):
            return True
        if re.fullmatch(r"^\d{1,2}/\d{1,2}(/\d{2,4})?$", t_clean):
            return True
        if len(t_clean) <= 1:
            return True
        return False

    tabela_linhas = []

    # A. Leitura se for XLSX (openpyxl)
    if conteudo_bytes and (nome_arquivo.endswith((".xlsx", ".xlsm", ".xltx")) or conteudo_bytes.startswith(b"PK\x03\x04")):
        try:
            import io
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(conteudo_bytes), data_only=True)
            for sheet in wb.worksheets:
                rows = list(sheet.iter_rows(values_only=True))
                for r in rows:
                    if r and any(c is not None and str(c).strip() for c in r):
                        tabela_linhas.append([str(c).strip() if c is not None else "" for c in r])
                if tabela_linhas:
                    break
        except Exception as e:
            print(f"[IMPORT_XLSX_ERROR] {e}")

    # B. Leitura se for XLS legado (xlrd)
    if not tabela_linhas and conteudo_bytes and (nome_arquivo.endswith(".xls") or conteudo_bytes.startswith(b"\xd0\xcf\x11\xe0")):
        try:
            import xlrd
            book = xlrd.open_workbook(file_contents=conteudo_bytes)
            for sheet_idx in range(book.nsheets):
                sh = book.sheet_by_index(sheet_idx)
                for rx in range(sh.nrows):
                    r_vals = [str(sh.cell_value(rx, cx)).strip() for cx in range(sh.ncols)]
                    if any(r_vals):
                        tabela_linhas.append(r_vals)
                if tabela_linhas:
                    break
        except Exception as e:
            print(f"[IMPORT_XLS_ERROR] {e}")

    # C. Leitura se for CSV ou texto puro
    if not tabela_linhas:
        if conteudo_bytes and not texto_plano:
            try:
                texto_plano = conteudo_bytes.decode("utf-8-sig", errors="ignore")
            except Exception:
                texto_plano = ""
        if texto_plano:
            import csv
            import io
            try:
                sniffer = csv.Sniffer()
                sample = texto_plano[:2048]
                dialect = sniffer.sniff(sample, delimiters=",;\t|")
                reader = csv.reader(io.StringIO(texto_plano), dialect)
                for row in reader:
                    if any(row):
                        tabela_linhas.append([str(c).strip() for c in row])
            except Exception:
                for line in texto_plano.splitlines():
                    if line.strip():
                        tabela_linhas.append([line.strip()])

    importacoes = []

    if tabela_linhas:
        # ESTRATÉGIA 1: Matriz com dias da semana nas colunas (Segunda, Terça, etc.)
        melhor_linha_dias = None
        col_dias = {}
        for r_idx in range(min(6, len(tabela_linhas))):
            dias_detectados = {}
            for c_idx, cell_val in enumerate(tabela_linhas[r_idx]):
                cell_ascii = unicodedata.normalize("NFKD", cell_val).encode("ascii", "ignore").decode("ascii").casefold()
                for chave, nome_oficial in dias_semana_oficiais:
                    padrao = rf"\b{chave}(?:-feira)?\b"
                    if chave == "segunda": padrao = r"\bseg(?:unda)?(?:-feira)?\b"
                    elif chave == "terca": padrao = r"\bter(?:ca|ça)?(?:-feira)?\b"
                    elif chave == "quarta": padrao = r"\bqua(?:rta)?(?:-feira)?\b"
                    elif chave == "quinta": padrao = r"\bqui(?:nta)?(?:-feira)?\b"
                    elif chave == "sexta": padrao = r"\bsex(?:ta)?(?:-feira)?\b"
                    elif chave == "sabado": padrao = r"\bs[aá]?b(?:ado)?\b"
                    elif chave == "domingo": padrao = r"\bdom(?:ingo)?\b"
                    if re.search(padrao, cell_ascii):
                        dias_detectados[c_idx] = nome_oficial
                        break
            if len(dias_detectados) >= 2:
                melhor_linha_dias = r_idx
                col_dias = dias_detectados
                break

        if col_dias and melhor_linha_dias is not None:
            for r_idx, row in enumerate(tabela_linhas[melhor_linha_dias + 1:]):
                hora_str = str(row[0] or "") if len(row) > 0 else ""
                hora_match = re.search(r"(\d{1,2}:\d{2})\s*(?:-|–|a|ate|até)\s*(\d{1,2}:\d{2})", hora_str, flags=re.IGNORECASE)
                hora_inicio = hora_match.group(1) if hora_match else None
                hora_fim = hora_match.group(2) if hora_match else None
                slot = r_idx if r_idx < len(default_slots) else 0
                h_ini = hora_inicio or default_slots[min(slot, len(default_slots)-1)][0]
                h_fim = hora_fim or default_slots[min(slot, len(default_slots)-1)][1]

                for c_idx, dia_nome in col_dias.items():
                    if c_idx < len(row):
                        mat = normalizar_materia(row[c_idx])
                        if mat and not eh_cabecalho_ou_invalido(mat):
                            importacoes.append({
                                "dia": dia_nome,
                                "slot": slot,
                                "materia": mat,
                                "horario_inicio": h_ini,
                                "horario_fim": h_fim,
                            })

        # ESTRATÉGIA 2: Tabela em colunas [Dia, Matéria]
        if not importacoes:
            col_dia_idx = None
            col_mat_idx = None
            for c_idx, cell_val in enumerate(tabela_linhas[0]):
                c_ascii = unicodedata.normalize("NFKD", cell_val).encode("ascii", "ignore").decode("ascii").casefold()
                if any(d in c_ascii for d in ["dia", "data", "semana"]):
                    col_dia_idx = c_idx
                elif any(m in c_ascii for m in ["materia", "disciplina", "assunto", "conteudo"]):
                    col_mat_idx = c_idx

            if col_dia_idx is not None and col_mat_idx is not None:
                contagem_por_dia = {}
                for row in tabela_linhas[1:]:
                    if len(row) > max(col_dia_idx, col_mat_idx):
                        dia_raw = unicodedata.normalize("NFKD", row[col_dia_idx]).encode("ascii", "ignore").decode("ascii").casefold()
                        mat = normalizar_materia(row[col_mat_idx])
                        if not mat or eh_cabecalho_ou_invalido(mat):
                            continue
                        dia_nome = None
                        for chave, oficial in dias_semana_oficiais:
                            if chave in dia_raw:
                                dia_nome = oficial
                                break
                        if not dia_nome:
                            dia_nome = "Segunda"
                        slot = contagem_por_dia.get(dia_nome, 0)
                        contagem_por_dia[dia_nome] = slot + 1
                        slot_clamped = min(slot, len(default_slots) - 1)
                        importacoes.append({
                            "dia": dia_nome,
                            "slot": slot_clamped,
                            "materia": mat,
                            "horario_inicio": default_slots[slot_clamped][0],
                            "horario_fim": default_slots[slot_clamped][1],
                        })

        # ESTRATÉGIA 3: UNIVERSAL (Aceita QUALQUER planilha Excel ou lista com matérias)
        if not importacoes:
            materias_encontradas = []
            for row in tabela_linhas:
                for cell in row:
                    mat = normalizar_materia(cell)
                    if mat and not eh_cabecalho_ou_invalido(mat):
                        if mat not in materias_encontradas:
                            materias_encontradas.append(mat)

            if materias_encontradas:
                dia_idx = 0
                slot_idx = 0
                for mat in materias_encontradas:
                    dia_nome = dias_ordenados[dia_idx % len(dias_ordenados)]
                    importacoes.append({
                        "dia": dia_nome,
                        "slot": slot_idx,
                        "materia": mat,
                        "horario_inicio": default_slots[slot_idx][0],
                        "horario_fim": default_slots[slot_idx][1],
                    })
                    dia_idx += 1
                    if dia_idx >= len(dias_ordenados):
                        dia_idx = 0
                        slot_idx = min(slot_idx + 1, len(default_slots) - 1)

    if not importacoes:
        return RedirectResponse(url=f"/estudos?success=erro_importar&semana={semana_inicio}", status_code=303)

    with Session(engine) as session:
        existentes = session.exec(select(EstudoPlano).where(
            EstudoPlano.usuario_id == usuario_id,
            EstudoPlano.semana_inicio == data_inicio,
        )).all()
        for plano in existentes:
            session.delete(plano)

        for item in importacoes:
            session.add(EstudoPlano(
                usuario_id=usuario_id,
                semana_inicio=data_inicio,
                materia=item["materia"],
                dias_semana=item["dia"],
                dia_semana=item["dia"],
                slot=item["slot"],
                horario_inicio=item["horario_inicio"],
                horario_fim=item["horario_fim"],
                meta_questoes=0,
            ))

        session.commit()

    return RedirectResponse(url=f"/estudos?success=plano_importado&semana={semana_inicio}", status_code=303)


@app.post("/estudos/grade", response_class=HTMLResponse)
async def save_estudo_grade(request: Request, semana_inicio: str = Form(...)):
    user_id = request.cookies.get("user_id")
    if not user_id:
        return RedirectResponse(url="/login")
    try:
        usuario_id = int(user_id)
        data_inicio = datetime.strptime(semana_inicio, "%Y-%m-%d").date()
        dias_grade = ["Segunda", "Terça", "Quarta", "Quinta", "Sexta", "Sábado", "Domingo"]
        horarios_grade = [("21:00", "22:00"), ("22:15", "00:15"), ("00:31", "01:15")]
        form = await request.form()
        for slot in range(len(horarios_grade)):
            inicio = str(form.get(f"horario_inicio_{slot}") or "").strip()
            fim = str(form.get(f"horario_fim_{slot}") or "").strip()
            if inicio and fim:
                horarios_grade[slot] = (inicio, fim)
        with Session(engine) as session:
            existentes = session.exec(select(EstudoPlano).where(
                EstudoPlano.usuario_id == usuario_id,
                EstudoPlano.semana_inicio == data_inicio,
            )).all()
            por_celula = {}
            for plano in existentes:
                if plano.slot is not None and plano.dia_semana:
                    por_celula.setdefault((plano.slot, plano.dia_semana), []).append(plano)

            for slot, (hora_inicio, hora_fim) in enumerate(horarios_grade):
                for dia_index, dia in enumerate(dias_grade):
                    campo = f"grade_{slot}_{dia_index}"
                    if campo not in form:
                        continue
                    materia = str(form.get(campo) or "").strip()
                    celula = por_celula.get((slot, dia), [])
                    if materia:
                        if celula:
                            for plano in celula:
                                plano.materia = materia
                                plano.dias_semana = dia
                                plano.horario_inicio = hora_inicio
                                plano.horario_fim = hora_fim
                        else:
                            session.add(EstudoPlano(
                                usuario_id=usuario_id,
                                semana_inicio=data_inicio,
                                materia=materia,
                                dias_semana=dia,
                                dia_semana=dia,
                                slot=slot,
                                horario_inicio=hora_inicio,
                                horario_fim=hora_fim,
                                meta_questoes=0,
                            ))
                    elif celula:
                        for plano in celula:
                            session.delete(plano)
            session.commit()
    except (ValueError, TypeError):
        return RedirectResponse(url="/estudos?success=erro", status_code=303)
    return RedirectResponse(url=f"/estudos?success=grade&semana={semana_inicio}", status_code=303)

@app.post("/estudos/plano", response_class=HTMLResponse)
async def save_estudo_plano(
    request: Request,
    semana_inicio: str = Form(...),
    materia: str = Form(...),
    dias_semana: str = Form(...),
    meta_questoes: int = Form(0),
):
    user_id = request.cookies.get("user_id")
    if not user_id:
        return RedirectResponse(url="/login")
    try:
        data_inicio = datetime.strptime(semana_inicio, "%Y-%m-%d").date()
        nome_materia = materia.strip()
        if not nome_materia:
            raise ValueError("matéria vazia")
        with Session(engine) as session:
            session.add(EstudoPlano(
                usuario_id=int(user_id),
                semana_inicio=data_inicio,
                materia=nome_materia,
                dias_semana=dias_semana.strip() or "Segunda",
                meta_questoes=max(0, meta_questoes),
            ))
            session.commit()
    except (ValueError, TypeError):
        return RedirectResponse(url="/estudos?success=erro", status_code=303)
    return RedirectResponse(url=f"/estudos?success=plano&semana={semana_inicio}", status_code=303)

@app.post("/estudos", response_class=HTMLResponse)
async def save_estudos(
    request: Request,
    registro_id: Optional[int] = Form(None),
    data_estudo: str = Form(...),
    horario_inicio: str = Form(""),
    horario_fim: str = Form(""),
    assunto: str = Form(""),
    tipo_registro: str = Form("estudo"),
    nome_simulado: str = Form(""),
    quantidade_questoes: int = Form(0),
    total_acertos: int = Form(0),
    total_erros: int = Form(0),
    observacao: str = Form(""),
    status: str = Form("concluido")
):
    user_id = request.cookies.get("user_id")
    if not user_id:
        return RedirectResponse(url="/login")

    try:
        with Session(engine) as session:
            sessao = session.get(EstudoSessao, registro_id) if registro_id else None
            if sessao and sessao.usuario_id != int(user_id):
                sessao = None
            if not sessao:
                sessao = EstudoSessao(usuario_id=int(user_id))
            sessao.data_estudo = datetime.strptime(data_estudo, "%Y-%m-%d").date()
            sessao.horario_inicio = horario_inicio or None
            sessao.horario_fim = horario_fim or None
            sessao.tipo_registro = "simulado" if tipo_registro == "simulado" else "estudo"
            assunto_limpo = (assunto or "").strip()
            sessao.assunto = assunto_limpo or ("Geral" if sessao.tipo_registro == "simulado" else "Estudo")
            sessao.nome_simulado = (nome_simulado.strip() or None) if sessao.tipo_registro == "simulado" else None
            questoes = max(0, quantidade_questoes)
            acertos = min(max(0, total_acertos), questoes)
            erros = max(0, total_erros)
            # Mantém o registro consistente: toda questão fica classificada
            # como acerto ou erro, sem permitir que os totais ultrapassem o total.
            if acertos + erros != questoes:
                erros = max(0, questoes - acertos)
            sessao.quantidade_questoes = questoes
            sessao.total_acertos = acertos
            sessao.total_erros = erros
            sessao.observacao = observacao.strip() or None
            sessao.status = status or "concluido"
            session.add(sessao)
            session.commit()
    except Exception as exc:
        return RedirectResponse(url="/estudos?success=erro#registro", status_code=303)

    return RedirectResponse(url="/estudos?success=ok#registro", status_code=303)

@app.post("/estudos/excluir")
async def delete_estudo(request: Request, registro_id: int = Form(...)):
    user_id = request.cookies.get("user_id")
    if not user_id:
        return RedirectResponse(url="/login")
    materia_reaberta = ""
    try:
        with Session(engine) as session:
            sessao = session.get(EstudoSessao, registro_id)
            if sessao and sessao.usuario_id == int(user_id):
                materia_reaberta = (sessao.assunto or "").strip()
                session.delete(sessao)
                session.commit()
    except (ValueError, TypeError):
        pass
    url = "/estudos?success=excluido"
    if materia_reaberta:
        url += f"&materia_reaberta={quote(materia_reaberta)}"
    return RedirectResponse(url=url, status_code=303)

@app.get("/cofre", response_class=HTMLResponse)
async def view_cofre(request: Request):
    user_id = request.cookies.get("user_id")
    user_role = get_signed_cookie(request, "user_role")
    if not user_id:
        return RedirectResponse(url="/login")
    if user_role != "admin":
        return RedirectResponse(url="/")
    return templates.TemplateResponse(
        request=request,
        name="cofre_senhas.html",
        context={
            "config": APP_CONFIG
        }
    )

@app.get("/configuracoes", response_class=HTMLResponse)
async def view_configuracoes(request: Request):
    user_id = request.cookies.get("user_id")
    user_role = request.cookies.get("user_role")
    user_perms = request.cookies.get("user_perms") or "[]"
    if not user_id: return RedirectResponse(url="/login")
    try:
        perms = json.loads(user_perms)
    except:
        perms = []
    if user_role != "admin" and "Ajustes" not in perms:
        return RedirectResponse(url="/")
    # Recarregar config para ter dados frescos
    fresh_config = load_app_config()
    return templates.TemplateResponse(
        request=request,
        name="configuracoes.html",
        context={"config": fresh_config, "message": "", "tipo_msg": ""}
    )

@app.post("/configuracoes", response_class=HTMLResponse)
async def save_configuracoes(
    request: Request,
    app_name: str = Form("PrintDash"),
    app_tagline: str = Form(""),
    app_color: str = Form("#6366f1"),
    app_price_bw: str = Form("0.50"),
    app_price_color: str = Form("1.00"),
    url_atendimento: str = Form(""),
    guia_suporte_url: str = Form(""),
    app_logo_url: str = Form(""),
    app_logo_remove: str = Form("0"),
    mod_impressao: Optional[str] = Form(None),
    mod_suporte: Optional[str] = Form(None),
    mod_atendimento: Optional[str] = Form(None),
    whatsapp_webhook_url: str = Form(""),
    whatsapp_webhook_token: str = Form(""),
    app_logo_upload: Optional[UploadFile] = File(None),
    guia_suporte_upload: Optional[UploadFile] = File(None)
):
    user_id = request.cookies.get("user_id")
    user_role = request.cookies.get("user_role")
    user_perms = request.cookies.get("user_perms") or "[]"
    if not user_id: return RedirectResponse(url="/login")
    try:
        perms = json.loads(user_perms)
    except:
        perms = []
    if user_role != "admin" and "Ajustes" not in perms:
        return RedirectResponse(url="/")
    
    import shutil
    message = ""
    tipo_msg = ""
    
    try:
        price_bw = float(str(app_price_bw).replace(",", "."))
        price_color = float(str(app_price_color).replace(",", "."))
        logo = app_logo_url
        guia_url = guia_suporte_url
        
        upload_dir = os.path.join(BASE_DIR, "public_html", "uploads")
        os.makedirs(upload_dir, exist_ok=True)

        # Remove somente os nomes de logo criados por este sistema.
        if app_logo_remove == "1":
            logo = ""
            for logo_ext in ("png", "jpg", "jpeg", "gif", "svg", "webp"):
                old_logo = os.path.join(upload_dir, f"logo_customizada.{logo_ext}")
                if os.path.isfile(old_logo):
                    os.remove(old_logo)
        
        # Upload de Logo
        if app_logo_upload and app_logo_upload.filename:
            ext = app_logo_upload.filename.split(".")[-1].lower()
            if ext in ["png", "jpg", "jpeg", "gif", "svg", "webp"]:
                new_name = f"logo_customizada.{ext}"
                file_path = os.path.join(upload_dir, new_name)
                with open(file_path, "wb") as buffer:
                    shutil.copyfileobj(app_logo_upload.file, buffer)
                logo = f"uploads/{new_name}?v={int(datetime.now().timestamp())}"
        
        # Upload de Guia
        if guia_suporte_upload and guia_suporte_upload.filename:
            ext = guia_suporte_upload.filename.split(".")[-1].lower()
            if ext in ["pdf", "doc", "docx", "txt", "jpg", "png"]:
                new_name = f"guia_suporte_usuario.{ext}"
                file_path = os.path.join(upload_dir, new_name)
                with open(file_path, "wb") as buffer:
                    shutil.copyfileobj(guia_suporte_upload.file, buffer)
                guia_url = f"uploads/{new_name}"
        
        new_config = {
            "APP_NAME": app_name,
            "APP_TAGLINE": app_tagline,
            "APP_LOGO_URL": logo,
            "APP_COLOR": app_color,
            "APP_PRICE_BW": price_bw,
            "APP_PRICE_COLOR": price_color,
            "MODULO_IMPRESSAO": mod_impressao is not None,
            "MODULO_SUPORTE": mod_suporte is not None,
            "MODULO_ATENDIMENTO": mod_atendimento is not None,
            "URL_ATENDIMENTO": url_atendimento,
            "GUIA_SUPORTE_URL": guia_url,
            "WHATSAPP_WEBHOOK_URL": whatsapp_webhook_url,
            "WHATSAPP_WEBHOOK_TOKEN": whatsapp_webhook_token
        }
        
        config_path = "public_html/includes/config.json"
        os.makedirs(os.path.dirname(config_path), exist_ok=True)
        with open(config_path, "w") as f:
            json.dump(new_config, f, indent=4)
        
        # Atualizar config global
        global APP_CONFIG
        APP_CONFIG = new_config
        
        message = "ConfiguraÃ§Ãµes atualizadas!"
        tipo_msg = "success"
    except Exception as e:
        message = f"Erro ao salvar: {str(e)}"
        tipo_msg = "error"
    
    fresh_config = load_app_config()
    return templates.TemplateResponse(
        request=request,
        name="configuracoes.html",
        context={"config": fresh_config, "message": message, "tipo_msg": tipo_msg}
    )



@app.get("/meu_painel", response_class=HTMLResponse)
async def view_meu_painel(request: Request):
    user_id = request.cookies.get("user_id")
    if not user_id:
        return RedirectResponse(url="/login")
        
    with Session(engine) as session:
        user = session.get(Usuario, int(user_id))
        if not user:
            return RedirectResponse(url="/login")
            
        agora = datetime.now()
        
        # MÃªs atual
        inicio_mes = datetime(agora.year, agora.month, 1)
        if agora.month == 12:
            fim_mes = datetime(agora.year + 1, 1, 1)
        else:
            fim_mes = datetime(agora.year, agora.month + 1, 1)
            
        # MÃªs anterior
        if agora.month == 1:
            inicio_mes_ant = datetime(agora.year - 1, 12, 1)
            fim_mes_ant = datetime(agora.year, 1, 1)
        else:
            inicio_mes_ant = datetime(agora.year, agora.month - 1, 1)
            fim_mes_ant = inicio_mes

        # Dados do usuÃ¡rio no mÃªs atual
        paginas_mes = session.exec(
            select(func.sum(Impressao.paginas))
            .where(Impressao.usuario == user.username)
            .where(Impressao.data_hora >= inicio_mes, Impressao.data_hora < fim_mes)
        ).one() or 0
        
        custo_mes = session.exec(
            select(func.sum(Impressao.valor_total))
            .where(Impressao.usuario == user.username)
            .where(Impressao.data_hora >= inicio_mes, Impressao.data_hora < fim_mes)
        ).one() or 0
        
        jobs_mes = session.exec(
            select(func.count(Impressao.id))
            .where(Impressao.usuario == user.username)
            .where(Impressao.data_hora >= inicio_mes, Impressao.data_hora < fim_mes)
        ).one() or 0

        # Dados do mÃªs anterior
        custo_mes_ant = session.exec(
            select(func.sum(Impressao.valor_total))
            .where(Impressao.usuario == user.username)
            .where(Impressao.data_hora >= inicio_mes_ant, Impressao.data_hora < fim_mes_ant)
        ).one() or 0
        
        # VariaÃ§Ã£o em %
        c_mes = float(custo_mes)
        c_ant = float(custo_mes_ant)
        if c_ant > 0:
            variacao_custo = ((c_mes - c_ant) / c_ant) * 100
        elif c_mes > 0:
            variacao_custo = 100.0
        else:
            variacao_custo = 0.0

        # Chamados recentes do usuÃ¡rio
        chamados_raw = session.exec(
            select(Chamado)
            .where(Chamado.usuario == user.username)
            .order_by(desc(Chamado.data_abertura))
            .limit(5)
        ).all()
        chamados = []
        for c in chamados_raw:
            d = c.dict()
            d["created_at"] = d.get("data_abertura")
            chamados.append(d)
        
        stats = {
            "paginas_mes": int(paginas_mes),
            "custo_mes": c_mes,
            "jobs_mes": int(jobs_mes),
            "variacao_custo": variacao_custo
        }

        return templates.TemplateResponse(
            request=request,
            name="meu_painel.html",
            context={
                "config": APP_CONFIG,
                "current_user": user,
                "stats": stats,
                "chamados": chamados
            }
        )

@app.get("/executivo", response_class=HTMLResponse)

async def view_executivo(request: Request):
    user_id = request.cookies.get("user_id")
    user_role = request.cookies.get("user_role")
    if not user_id:
        return RedirectResponse(url="/login")
    if user_role not in ("admin", "gestor"):
        return RedirectResponse(url="/")

    with Session(engine) as session:
        agora = datetime.now()
        hoje = agora.date()

        # --- MÃŠS ATUAL ---
        inicio_mes = datetime(agora.year, agora.month, 1)
        if agora.month == 12:
            fim_mes = datetime(agora.year + 1, 1, 1)
        else:
            fim_mes = datetime(agora.year, agora.month + 1, 1)

        paginas_mes = session.exec(
            select(func.sum(Impressao.paginas)).where(
                Impressao.data_hora >= inicio_mes, Impressao.data_hora < fim_mes
            )
        ).one() or 0
        custo_mes = session.exec(
            select(func.sum(Impressao.valor_total)).where(
                Impressao.data_hora >= inicio_mes, Impressao.data_hora < fim_mes
            )
        ).one() or 0
        custo_mes = float(custo_mes)

        # --- PROJEÃ‡ÃƒO ATÃ‰ FIM DO MÃŠS ---
        dia_atual = agora.day
        total_dias_mes = (fim_mes - inicio_mes).days
        custo_projetado = (custo_mes / dia_atual * total_dias_mes) if dia_atual > 0 else 0
        progresso_mes = round(dia_atual / total_dias_mes * 100)

        # --- MÃŠS ANTERIOR ---
        if agora.month == 1:
            inicio_mes_ant = datetime(agora.year - 1, 12, 1)
            fim_mes_ant = datetime(agora.year, 1, 1)
        else:
            inicio_mes_ant = datetime(agora.year, agora.month - 1, 1)
            fim_mes_ant = inicio_mes

        custo_mes_ant = session.exec(
            select(func.sum(Impressao.valor_total)).where(
                Impressao.data_hora >= inicio_mes_ant, Impressao.data_hora < fim_mes_ant
            )
        ).one() or 0
        custo_mes_ant = float(custo_mes_ant)

        paginas_mes_ant = session.exec(
            select(func.sum(Impressao.paginas)).where(
                Impressao.data_hora >= inicio_mes_ant, Impressao.data_hora < fim_mes_ant
            )
        ).one() or 0

        # --- VARIAÃ‡ÃƒO % ---
        if custo_mes_ant > 0:
            variacao_pct = round((custo_projetado - custo_mes_ant) / custo_mes_ant * 100, 1)
        else:
            variacao_pct = 0

        # --- TOTAL GERAL HISTÃ“RICO ---
        custo_total = float(session.exec(select(func.sum(Impressao.valor_total))).one() or 0)
        paginas_total = int(session.exec(select(func.sum(Impressao.paginas))).one() or 0)

        # --- TOP 8 USUÃRIOS DO MÃŠS ---
        top_usuarios_raw = session.exec(
            select(
                Impressao.usuario,
                func.sum(Impressao.paginas).label("paginas"),
                func.sum(Impressao.valor_total).label("custo"),
                func.count(Impressao.id).label("jobs")
            )
            .where(Impressao.data_hora >= inicio_mes, Impressao.data_hora < fim_mes)
            .group_by(Impressao.usuario)
            .order_by(func.sum(Impressao.valor_total).desc())
            .limit(8)
        ).all()
        top_usuarios = [
            {"nome": r[0] or "Desconhecido", "paginas": int(r[1] or 0),
             "custo": float(r[2] or 0), "jobs": int(r[3] or 0)}
            for r in top_usuarios_raw
        ]

        # --- EVOLUÃ‡ÃƒO DIÃRIA (Ãºltimos 30 dias) ---
        inicio_30 = agora - timedelta(days=30)
        evol_raw = session.exec(
            select(
                func.date(Impressao.data_hora).label("dia"),
                func.sum(Impressao.valor_total).label("custo")
            )
            .where(Impressao.data_hora >= inicio_30)
            .group_by(func.date(Impressao.data_hora))
            .order_by(func.date(Impressao.data_hora))
        ).all()
        evolucao = []
        for row in evol_raw:
            try:
                d = datetime.strptime(str(row[0]), "%Y-%m-%d")
                evolucao.append({"dia": d.strftime("%d/%m"), "custo": round(float(row[1] or 0), 2)})
            except:
                pass

        # --- RANKING POR TIPO DE IMPRESSÃƒO ---
        tipo_raw = session.exec(
            select(
                Impressao.tipo_impressao,
                func.sum(Impressao.paginas).label("paginas"),
                func.sum(Impressao.valor_total).label("custo")
            )
            .where(Impressao.data_hora >= inicio_mes, Impressao.data_hora < fim_mes)
            .group_by(Impressao.tipo_impressao)
        ).all()
        por_tipo = [
            {"tipo": r[0] or "Outro", "paginas": int(r[1] or 0), "custo": float(r[2] or 0)}
            for r in tipo_raw
        ]

        # --- NOME DO MÃŠS ATUAL E ANTERIOR ---
        MESES_PT = ["Janeiro","Fevereiro","MarÃ§o","Abril","Maio","Junho",
                    "Julho","Agosto","Setembro","Outubro","Novembro","Dezembro"]
        nome_mes = MESES_PT[agora.month - 1]
        nome_mes_ant = MESES_PT[inicio_mes_ant.month - 1]

        return templates.TemplateResponse(
            request=request,
            name="executivo.html",
            context={
                "config": APP_CONFIG,
                "nome_mes": nome_mes,
                "nome_mes_ant": nome_mes_ant,
                "ano": agora.year,
                "custo_mes": custo_mes,
                "paginas_mes": int(paginas_mes),
                "custo_projetado": round(custo_projetado, 2),
                "custo_mes_ant": custo_mes_ant,
                "paginas_mes_ant": int(paginas_mes_ant),
                "variacao_pct": variacao_pct,
                "progresso_mes": progresso_mes,
                "custo_total": custo_total,
                "paginas_total": paginas_total,
                "top_usuarios": top_usuarios,
                "evolucao": evolucao,
                "por_tipo": por_tipo,
                "hoje": hoje.strftime("%d/%m/%Y"),
            }
        )


@app.get("/login", response_class=HTMLResponse)
async def view_login(request: Request, intent: str = ""):
    return templates.TemplateResponse(
        request=request,
        name="login.html",
        context={"config": APP_CONFIG, "intent": intent, "only_atendimento": False}
    )

@app.get("/estudos/login", response_class=HTMLResponse)
async def view_estudos_login(request: Request):
    return templates.TemplateResponse(
        request=request,
        name="estudos_login.html",
        context={"config": APP_CONFIG},
    )

@app.get("/atendimento/login", response_class=HTMLResponse)
async def view_atendimento_login(request: Request):
    return templates.TemplateResponse(
        request=request,
        name="atendimento_login.html",
        context={"config": APP_CONFIG}
    )

def _can_access_atendimento(request: Request) -> bool:
    """Permite a Central de Atendimento a administradores ou usuários autorizados."""
    if get_signed_cookie(request, "user_role") == "admin":
        return True
    if get_signed_cookie(request, "user_role") in ("recepcao", "recepção", "recepcionista"):
        return True
    if get_signed_cookie(request, "auth_user_role") in ("recepcao", "recepção", "recepcionista"):
        return True
    try:
        raw_perms = get_signed_cookie(request, "user_perms", "[]") or "[]"
        perms = normalize_permissions(json.loads(raw_perms))
        return "Atendimento" in perms or "Suporte" in perms
    except Exception:
        return False


@app.get("/central-atendimento", response_class=HTMLResponse)
async def view_central_atendimento(request: Request):
    desktop_mode = request.query_params.get("desktop") == "1" or request.cookies.get("central_atendimento_tauri") == "1"
    user_id = request.cookies.get("user_id")
    user_role = request.cookies.get("user_role")
    user_name = request.cookies.get("user_name")
    profile_name = user_name or ""
    profile_photo_url = ""

    config = load_app_config()
    try:
        with Session(engine) as session:
            current_user = session.get(Usuario, int(user_id)) if user_id else None
        if current_user:
            profile_name = current_user.nome or current_user.username or profile_name
            profile_photo_url = current_user.foto_perfil_url or ""
    except Exception:
        pass
    if user_role == "cliente_atendimento":
        try:
            with Session(engine) as session:
                cliente = session.get(Usuario, int(user_id))
            if cliente:
                config = dict(config)
                config["APP_NAME"] = cliente.empresa_nome or cliente.nome or config.get("APP_NAME", "Central de Atendimento")
                config["APP_LOGO_URL"] = cliente.empresa_logo_url or ""
        except Exception as exc:
            print(f"[CENTRAL] Não foi possível carregar a marca do cliente: {exc}")
    if not config.get("MODULO_ATENDIMENTO", False):
        return RedirectResponse(url="/")

    if not user_id:
        return templates.TemplateResponse(
            request=request,
            name="atendimento_login.html",
            context={"config": config},
        )

    has_access = _can_access_atendimento(request)
    if not has_access:
        return RedirectResponse(url="/")
    can_edit_identity = api_whatsapp._can_edit_identity(request)
    can_manage_company = user_role == "admin" or (user_role == "cliente_atendimento" and can_edit_identity)

    response = templates.TemplateResponse(
        request=request,
        name="atendimento_whatsapp.html",
        context={
            "config": config,
            "is_admin": user_role == "admin",
            "user_role": user_role,
            "can_manage_atendimento": user_role in ("admin", "gestor", "gerente", "supervisor", "recepcao", "recepção", "recepcionista"),
            "is_reception": user_role in ("recepcao", "recepção", "recepcionista"),
            "user_name": user_name,
            "authenticated": True,
            "has_access": has_access,
            "desktop_mode": desktop_mode,
            "can_edit_identity": user_role in ("admin", "operator", "gestor", "gerente", "supervisor", "recepcao", "recepÃ§Ã£o", "recepcionista"),
            "profile_name": profile_name,
            "profile_photo_url": profile_photo_url,
            "can_edit_identity": can_edit_identity,
            "can_manage_company": can_manage_company,
            "atendimento_multiempresa_enabled": os.getenv("ATENDIMENTO_MULTIEMPRESA_ENABLED", "false").lower() in ("1", "true", "yes", "on"),
        }
    )
    response.headers["Cache-Control"] = "no-store, no-cache, must-revalidate, max-age=0"
    response.headers["Pragma"] = "no-cache"
    response.headers["Expires"] = "0"
    if request.query_params.get("desktop") == "1":
        response.set_cookie("central_atendimento_tauri", "1", httponly=False, samesite="lax")
    return response

@app.get("/atendimento")
async def redirect_atendimento():
    return RedirectResponse(url="/central-atendimento")

@app.get("/central-atendimento/dashboard", response_class=HTMLResponse)
async def view_central_atendimento_dashboard(request: Request):
    if request.query_params.get("desktop") == "1":
        return RedirectResponse(url="/central-atendimento?desktop=1", status_code=303)
    user_id = request.cookies.get("user_id")
    user_role = request.cookies.get("user_role")
    if not user_id:
        return RedirectResponse(url="/login")
    if user_role not in ("admin", "recepcao", "recepção", "recepcionista"):
        return RedirectResponse(url="/central-atendimento")

    config = load_app_config()
    with Session(engine) as session:
        chamados = session.exec(select(Chamado).where(Chamado.origem == "WhatsApp")).all()
        try:
            status_usuarios = session.connection().execute(text("SELECT status_conta, last_seen FROM usuarios")).fetchall()
            online_cutoff = datetime.now(timezone.utc) - timedelta(minutes=2)
            usuarios_ativos = sum(1 for row in status_usuarios if row[1] and row[1] >= online_cutoff)
            usuarios_inativos = len(status_usuarios) - usuarios_ativos
            total_usuarios = len(status_usuarios)
        except Exception as exc:
            print(f"[CENTRAL DASHBOARD] Nao foi possivel carregar status dos usuarios: {exc}")
            usuarios_ativos = 0
            usuarios_inativos = 0
            total_usuarios = 0
        operator_counts = {}
        atendimentos = [
            item for item in chamados
            if not str(item.whatsapp_cliente or "").lower().endswith("@g.us")
        ]
        for item in chamados:
            if not str(item.whatsapp_cliente or "").lower().endswith("@g.us") and item.assigned_user and item.status != "Resolvido":
                operator_counts[item.assigned_user] = operator_counts.get(item.assigned_user, 0) + 1
        def data_abertura_utc(item):
            valor = item.data_abertura
            if not valor:
                return datetime.min.replace(tzinfo=timezone.utc)
            return valor.replace(tzinfo=timezone.utc) if valor.tzinfo is None else valor.astimezone(timezone.utc)

        recent = sorted(atendimentos, key=data_abertura_utc, reverse=True)[:8]
        agora = datetime.now(timezone.utc)
        def idade_em_minutos(item):
            if not item.data_abertura:
                return 0
            return max(0, int((agora - data_abertura_utc(item)).total_seconds() // 60))
        # Grupos ficam abertos e não precisam ser assumidos; portanto não entram na triagem.
        triagem_aberta = [
            item for item in chamados
            if item.status != "Resolvido"
            and not item.assigned_user
            and not str(item.whatsapp_cliente or "").lower().endswith("@g.us")
        ]
        tempos_triagem = [idade_em_minutos(item) for item in triagem_aberta if item.data_abertura]
        triagem_urgente = sorted(
            [item for item in triagem_aberta if item.data_abertura and idade_em_minutos(item) >= 15],
            key=data_abertura_utc,
        )[:5]
        metrics = {
            "total": len(atendimentos),
            "triagem": len(triagem_aberta),
            "atendimento": sum(1 for item in atendimentos if item.status == "Em Atendimento"),
            "finalizados": sum(1 for item in atendimentos if item.status == "Resolvido"),
            "grupos": sum(1 for item in chamados if str(item.whatsapp_cliente or "").lower().endswith("@g.us")),
            "grupos_abertos": sum(1 for item in chamados if item.status != "Resolvido" and str(item.whatsapp_cliente or "").lower().endswith("@g.us")),
            "nao_lidas": sum(int(item.unread_admin or 0) for item in atendimentos),
            "usuarios_ativos": usuarios_ativos,
            "usuarios_inativos": usuarios_inativos,
            "total_usuarios": total_usuarios,
            "usuarios_ativos_percentual": round(usuarios_ativos / total_usuarios * 100) if total_usuarios else 0,
            "tempo_medio_espera": round(sum(tempos_triagem) / len(tempos_triagem)) if tempos_triagem else 0,
            "triagem_urgente": [
                {"id": item.id, "name": item.usuario or "Sem nome", "minutes": idade_em_minutos(item)}
                for item in triagem_urgente
            ],
            "operators": sorted(operator_counts.items(), key=lambda pair: pair[1], reverse=True),
            "recent": [
                {"id": item.id, "name": item.usuario or "Sem nome", "status": item.status, "assigned": item.assigned_user or "Sem operador", "date": item.data_abertura.strftime("%d/%m/%Y %H:%M")}
                for item in recent
            ],
        }
    return templates.TemplateResponse(
        request=request,
        name="atendimento_dashboard.html",
        context={"config": config, "metrics": metrics, "authenticated": True, "user_role": user_role, "user_name": request.cookies.get("user_name", "")},
    )

@app.get("/central-atendimento/contatos", response_class=HTMLResponse)
async def view_whatsapp_contacts(request: Request):
    desktop_mode = request.query_params.get("desktop") == "1" or request.cookies.get("central_atendimento_tauri") == "1"
    if not desktop_mode and "desktop=1" in (request.headers.get("referer") or ""):
        return RedirectResponse(url="/central-atendimento/contatos?desktop=1", status_code=307)

    user_id = request.cookies.get("user_id")
    user_role = request.cookies.get("user_role")
    if not user_id:
        return RedirectResponse(url="/login")
    if not _can_access_atendimento(request):
        return RedirectResponse(url="/central-atendimento")

    response = templates.TemplateResponse(
        request=request,
        name="whatsapp_contacts.html",
        context={
            "config": load_app_config(),
            "is_admin": user_role == "admin",
            "user_role": user_role,
            "authenticated": True,
            "user_name": request.cookies.get("user_name", ""),
            "desktop_mode": desktop_mode,
        },
    )
    if request.query_params.get("desktop") == "1":
        response.set_cookie("central_atendimento_tauri", "1", httponly=False, samesite="lax")
    return response

@app.get("/central-atendimento/configuracoes", response_class=HTMLResponse)
async def view_whatsapp_config(request: Request):
    desktop_mode = request.query_params.get("desktop") == "1" or request.cookies.get("central_atendimento_tauri") == "1"
    if not desktop_mode and "desktop=1" in (request.headers.get("referer") or ""):
        return RedirectResponse(url="/central-atendimento/configuracoes?desktop=1", status_code=307)
    user_id = request.cookies.get("user_id")
    user_role = request.cookies.get("user_role")
    if not user_id:
        return RedirectResponse(url="/login")

    allowed_roles = ("admin", "operator", "gestor", "gerente", "supervisor", "recepcao", "recepÃ§Ã£o", "recepcionista")
    tenant_admin_identity = api_whatsapp._can_edit_identity(request)
    if user_role not in allowed_roles and not tenant_admin_identity:
        return RedirectResponse(url="/central-atendimento")

    config = load_app_config()
    with Session(engine) as session:
        current_user = session.get(Usuario, int(user_id))
    return templates.TemplateResponse(
        request=request,
        name="whatsapp_config.html",
        context={
            "config": config,
            "is_admin": user_role == "admin",
            "authenticated": True,
            "user_name": request.cookies.get("user_name", ""),
            "desktop_mode": desktop_mode,
            "can_edit_identity": user_role in allowed_roles or tenant_admin_identity,
            "profile_name": (current_user.nome or current_user.username) if current_user else "",
            "profile_photo_url": current_user.foto_perfil_url if current_user else "",
        }
    )

@app.get("/central-atendimento/grupos", response_class=HTMLResponse)
async def view_whatsapp_groups(request: Request):
    user_id = request.cookies.get("user_id")
    if not user_id:
        return RedirectResponse(url="/login")
    config = load_app_config()
    groups = []
    with Session(engine) as session:
        chamados = session.exec(select(Chamado).where(Chamado.origem == "WhatsApp", Chamado.whatsapp_cliente.like("%@g.us")).order_by(desc(Chamado.data_abertura))).all()
        seen = set()
        for chamado in chamados:
            if chamado.whatsapp_cliente in seen:
                continue
            seen.add(chamado.whatsapp_cliente)
            group_name = chamado.usuario or ""
            if group_name in ("", "Grupo WhatsApp", chamado.whatsapp_cliente):
                group_name = api_whatsapp.get_whatsapp_group_name(chamado.whatsapp_cliente, chamado.whatsapp_instance)
                if group_name:
                    chamado.usuario = group_name
                    session.add(chamado)
                    session.commit()
            groups.append({"name": group_name or "Grupo WhatsApp", "jid": chamado.whatsapp_cliente, "last_activity": chamado.data_abertura.strftime("%d/%m/%Y %H:%M")})
    return templates.TemplateResponse(request=request, name="whatsapp_groups.html", context={"config": config, "groups": groups, "authenticated": True, "user_name": request.cookies.get("user_name", "")})

@app.get("/whatsapp_config")
async def redirect_whatsapp_config():
    return RedirectResponse(url="/central-atendimento/configuracoes")


@app.get("/central-atendimento/empresas", response_class=HTMLResponse)
async def view_atendimento_companies(request: Request):
    if not get_signed_cookie(request, "user_id"):
        return RedirectResponse(url="/login")
    is_platform_admin = get_signed_cookie(request, "user_role") == "admin"
    if not is_platform_admin:
        try:
            tenant_context = resolve_context(request)
            with Session(engine) as session:
                is_tenant_admin = bool(session.exec(text("""
                    SELECT 1 FROM usuarios_empresas_atendimento
                    WHERE empresa_atendimento_id = :tenant_id AND usuario_id = :user_id
                      AND papel = 'tenant_admin' AND ativo = TRUE
                """).bindparams(tenant_id=tenant_context.tenant_id, user_id=int(tenant_context.user_id))).first())
        except Exception:
            is_tenant_admin = False
        if not is_tenant_admin:
            return RedirectResponse(url="/central-atendimento")
    return templates.TemplateResponse(
        request=request,
        name="atendimento_empresas.html",
        context={"config": load_app_config(), "user_name": get_signed_cookie(request, "user_name", ""), "is_platform_admin": is_platform_admin},
    )
